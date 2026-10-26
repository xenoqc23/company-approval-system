"""New-request badges and trimmed Excel, on a fresh fictional local instance."""
import argparse,json
from pathlib import Path
from urllib.parse import urlsplit
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright,expect
parser=argparse.ArgumentParser();parser.add_argument('--base-url',default='http://127.0.0.1:18161');parser.add_argument('--password-file',type=Path,required=True);parser.add_argument('--output',type=Path,required=True);args=parser.parse_args()
base=args.base_url.rstrip('/');assert urlsplit(base).hostname in ['localhost','127.0.0.1']
args.output.mkdir(parents=True,exist_ok=True);password=args.password_file.read_text().strip();errors=[]
with sync_playwright() as p:
 browser=p.chromium.launch(executable_path='/usr/bin/chromium',args=['--no-sandbox']);pages={}
 def role(who):
  if who not in pages:
   context=browser.new_context(viewport={'width':1440,'height':1000},locale='ko-KR',timezone_id='Asia/Seoul');page=context.new_page();page.on('pageerror',lambda e:errors.append(str(e)));page.on('dialog',lambda d:d.accept())
   page.goto(base+'/login/');page.locator('[name=username]').fill('p_'+who);page.locator('[name=password]').fill(password);page.get_by_role('button',name='로그인',exact=True).click();page.wait_for_url(base+'/');expect(page.locator('#integrated-role')).to_be_visible();pages[who]=page
  return pages[who]
 def click(page,label):page.get_by_role('button',name=label,exact=True).click();page.wait_for_load_state('networkidle')
 def send(names):
  owner=role('dev_manager');checker=role('dev_director');ids=[]
  for name in names:
   owner.goto(base+'/purchases/new/');owner.locator('[name=request_method][value=direct]').check()
   for field,value in dict(product=name,vendor='가상구매처',quantity='2',unit_price='50000',url='https://example.com/equipment',reason='업무용 장비 교체').items():owner.locator('#id_'+field).fill(value)
   click(owner,'구매담당자에게 요청');owner.wait_for_url('**/requests/*/');pk=urlsplit(owner.url).path.rstrip('/').split('/')[-1];ids.append(pk)
   checker.goto(base+f'/purchases/requests/{pk}/');click(checker,'구매 내용 확인')
  finance=role('finance');finance.goto(base+'/purchases/reports/?tab=items')
  for pk in ids:finance.locator(f'[name=requests][value="{pk}"]').check()
  click(finance,'선택 품목으로 결제 요청');finance.locator('[name=note]').fill('가상 장비 결제 요청');click(finance,'대표님께 결제 요청');finance.wait_for_url('**/reports/*/');return urlsplit(finance.url).path
 def refresh_home(ceo,count):
  ceo.bring_to_front();ceo.evaluate("window.dispatchEvent(new Event('xenosys-refresh-notifications'))")
  expect(ceo.locator('main [data-live-purchase=new_purchase]')).to_have_text(str(count),timeout=20000)
  panel=ceo.locator('main [data-live-purchase-region=new_purchase]');badge=ceo.locator('#workspace-navigation [data-live-purchase-region=new_purchase]')
  if count:expect(panel).to_be_visible();expect(badge).to_be_visible()
  else:expect(panel).to_be_hidden();expect(badge).to_be_hidden()
  assert ceo.request.get(base+'/notifications/status/').json()['waiting']==0
 def export(page,label):
  with page.expect_download() as event:click(page,'엑셀 다운로드')
  path=args.output/(label+'.xlsx');event.value.save_as(path);sheet=load_workbook(path).active
  assert sheet.max_column==10;assert sheet.cell(1,10).value=='구매사유';assert sheet.cell(2,10).value=='업무용 장비 교체'
  assert sheet.cell(2,9).value==100000
  assert not {'진행상황','사유','수령 상태'}.intersection(c.value for c in sheet[1])
 ceo=role('ceo');refresh_home(ceo,0)
 first=send(['새요청 시험 모니터','새요청 시험 키보드']);refresh_home(ceo,1)
 second=send(['새요청 시험 장비']);refresh_home(ceo,2);ceo.screenshot(path=str(args.output/'01-home-new.png'),full_page=True)
 finance=role('finance');finance.goto(base+first);export(finance,'finance-export');refresh_home(ceo,2)
 detail=ceo.context.new_page();detail.on('pageerror',lambda e:errors.append(str(e)))
 detail.goto(base+'/purchases/reports/?new=1');assert detail.locator('tbody tr').count()==2
 detail.screenshot(path=str(args.output/'02-new-list.png'),full_page=True)
 detail.goto(base+first);expect(detail.locator('.purchase-content .heading')).to_contain_text('결제 대기');assert detail.locator('[data-payment-form]').count()==0
 assert detail.get_by_role('button',name='구매 승인',exact=True).count()==0
 export(detail,'ceo-export');refresh_home(ceo,1)
 detail.goto(base+'/purchases/reports/?new=1');assert detail.locator('tbody tr').count()==1
 detail.goto(base+second);refresh_home(ceo,0)
 detail.goto(base+'/purchases/reports/');assert detail.locator('tbody tr').count()==2
 detail.screenshot(path=str(args.output/'03-viewed-list.png'),full_page=True)
 # Opening requests never completes them; accounting still registers actual payment.
 finance.goto(base+first);expect(finance.locator('tbody tr').first).to_contain_text('결제 대기');finance.locator('[name=items]').first.check();click(finance,'선택 품목 결제 완료');click(finance,'결제 완료 확정')
 expect(finance.locator('tbody tr').first).to_contain_text('결제 완료');expect(finance.locator('tbody tr').nth(1)).to_contain_text('결제 대기')
 for page,path in [(ceo,'/'),(detail,'/purchases/reports/?new=1'),(detail,first)]:
  for width in [320,390,768,1440]:
   page.set_viewport_size({'width':width,'height':1000});assert page.goto(base+path).status==200
   assert page.evaluate('document.documentElement.scrollWidth <= innerWidth + 1'),(path,width)
 assert not errors,errors
 result=dict(checks=['live new-request count and zero visibility','counts batches not items','finance/export do not clear new','CEO detail clears only viewed request','new-only filter and original list','approval deferred and payment unchanged','10-column Excel preserves purchase purpose and prices','320–1440px no overflow'],js_errors=errors)
 (args.output/'result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False));browser.close()
