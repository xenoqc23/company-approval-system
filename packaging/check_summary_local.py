"""Browser check for the separately seeded fictional accounting summary."""
import argparse,json
from pathlib import Path
from urllib.parse import urlsplit
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright,expect
parser=argparse.ArgumentParser();parser.add_argument('--base-url',default='http://127.0.0.1:18171');parser.add_argument('--password-file',type=Path,required=True);parser.add_argument('--ids-file',type=Path,required=True);parser.add_argument('--dates-file',type=Path,required=True);parser.add_argument('--output',type=Path,required=True);args=parser.parse_args()
base=args.base_url.rstrip('/');assert urlsplit(base).hostname in ['localhost','127.0.0.1'];args.output.mkdir(parents=True,exist_ok=True)
ids=json.loads(args.ids_file.read_text());dates=json.loads(args.dates_file.read_text());errors=[]
with sync_playwright() as p:
 browser=p.chromium.launch(executable_path='/usr/bin/chromium',args=['--no-sandbox']);context=browser.new_context(viewport={'width':1440,'height':1000},locale='ko-KR',timezone_id='Asia/Seoul');page=context.new_page();page.on('pageerror',lambda error:errors.append(str(error)))
 page.goto(base+'/login/');page.locator('[name=username]').fill('p_finance');page.locator('[name=password]').fill(args.password_file.read_text().strip());page.get_by_role('button',name='로그인',exact=True).click();page.wait_for_url(base+'/');expect(page.locator('#integrated-role')).to_be_visible()
 page.goto(base+'/purchases/reports/');page.get_by_role('link',name='구매 집계',exact=True).click()
 expect(page.locator('[data-summary-total]')).to_have_text('450,000원');expect(page.locator('[data-summary-count]')).to_have_text('2건')
 expect(page.locator('.summary-detail-table tbody')).not_to_contain_text('집계 제외 결제 대기');page.screenshot(path=str(args.output/'01-this-month.png'),full_page=True)
 def apply():page.get_by_role('button',name='집계 보기',exact=True).click();page.wait_for_load_state('networkidle')
 def person(who):return page.locator(f'[name=employees][value="{ids["p_"+who]}"]')
 person('dev_manager').check();person('quality_lead').check();apply();expect(page.locator('[data-summary-total]')).to_have_text('300,000원')
 page.locator('[name=start]').fill(dates['prior']);page.locator('[name=end]').fill(dates['today']);apply();expect(page.locator('[data-summary-total]')).to_have_text('400,000원');expect(page.locator('[data-summary-count]')).to_have_text('2건')
 page.locator('[name=department]').select_option('품질팀');expect(person('dev_manager')).to_be_disabled();person('quality_manager').check();apply()
 expect(page.locator('[data-summary-total]')).to_have_text('250,000원');expect(page.locator('[data-summary-count]')).to_have_text('2건')
 expect(page.locator('.summary-detail-table')).not_to_contain_text('집계용 모니터');expect(page.locator('.summary-detail-table')).to_contain_text('집계용 허브');page.screenshot(path=str(args.output/'02-selected-people.png'),full_page=True)
 with page.expect_download() as event:page.get_by_role('link',name='현재 조건으로 엑셀 다운로드',exact=True).click()
 excel=args.output/'selected-summary.xlsx';event.value.save_as(excel);book=load_workbook(excel)
 assert book.sheetnames==['집계 요약','구매내역'];assert book['집계 요약']['B5'].value==250000
 detail=book['구매내역'];assert detail.max_row==3;assert detail.max_column==9
 assert {detail.cell(row,4).value for row in [2,3]}=={'집계용 키보드','집계용 허브'}
 page.get_by_role('button',name='전체 직원',exact=True).click();page.locator('[name=department]').select_option('');apply()
 expect(page.locator('[data-summary-total]')).to_have_text('550,000원');expect(page.locator('[data-summary-count]')).to_have_text('3건')
 page.get_by_role('button',name='이번 달',exact=True).click();apply();expect(page.locator('[data-summary-total]')).to_have_text('450,000원')
 for width in [320,390,768,1440]:
  page.set_viewport_size({'width':width,'height':1000});page.goto(base+'/purchases/summary/')
  assert page.evaluate('document.documentElement.scrollWidth <= innerWidth + 1'),width
 page.set_viewport_size({'width':390,'height':1000});page.screenshot(path=str(args.output/'03-mobile.png'),full_page=True)
 page.locator('[name=start]').fill('2026-12-01');page.locator('[name=end]').fill('2026-11-01');apply();expect(page.locator('[role=alert]')).to_contain_text('시작일')
 assert page.locator('[data-summary-total]').count()==0;assert page.get_by_role('link',name='현재 조건으로 엑셀 다운로드',exact=True).count()==0
 assert not errors,errors
 result=dict(checks=['default month totals','multiple employees across teams','inclusive period change','team filters reset hidden employees','received purchases included and pending excluded','filtered two-sheet Excel','clear selections means whole team','320–1440px layout','invalid filters do not show misleading totals'],js_errors=errors)
 (args.output/'result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False));browser.close()
