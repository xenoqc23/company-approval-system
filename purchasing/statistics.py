"""Read-only purchase totals from completed payments, within accounting access."""
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from decimal import Decimal
import json

from django import forms
from django.utils import timezone


class SummaryFilterForm(forms.Form):
    start = forms.DateField(label='시작일', widget=forms.DateInput(attrs={'type': 'date'}))
    end = forms.DateField(label='종료일', widget=forms.DateInput(attrs={'type': 'date'}))
    department = forms.ChoiceField(label='팀', required=False)
    employees = forms.MultipleChoiceField(label='직원', required=False)

    def __init__(self, *args, scope, **kwargs):
        super().__init__(*args, **kwargs)
        people = {}
        departments = set()
        for record in scope.order_by().values('owner_id', 'owner__first_name', 'owner__username', 'department').distinct():
            department = record['department']
            departments.add(department)
            person = people.setdefault(str(record['owner_id']), {
                'id': str(record['owner_id']), 'name': record['owner__first_name'] or record['owner__username'], 'departments': set(),
            })
            person['departments'].add(department)
        self.people = sorted(people.values(), key=lambda person: (person['name'], int(person['id'])))
        self.fields['department'].choices = [('', '전체 팀')] + [(value, value) for value in sorted(departments)]
        self.fields['employees'].choices = [(person['id'], person['name']) for person in self.people]
        selected = set(self.data.getlist('employees')) if hasattr(self.data, 'getlist') else set(self.data.get('employees', []))
        for person in self.people:
            person['selected'] = person['id'] in selected
            person['department_label'] = ' · '.join(sorted(person['departments']))
            person['departments_json'] = json.dumps(sorted(person['departments']), ensure_ascii=False)

    def clean(self):
        values = super().clean()
        start, end = values.get('start'), values.get('end')
        if start and end and start > end:
            raise forms.ValidationError('시작일은 종료일보다 늦을 수 없습니다.')
        if start and start < date(1900, 1, 1) or end and end > date(9998, 12, 31):
            raise forms.ValidationError('집계 기간을 확인해 주세요.')
        department = values.get('department')
        selected = set(values.get('employees', []))
        if department and any(person['id'] in selected and department not in person['departments'] for person in self.people):
            raise forms.ValidationError('선택한 팀에 해당하는 직원을 선택해 주세요.')
        return values


def aggregate(scope, values):
    tz = timezone.get_current_timezone()
    start = timezone.make_aware(datetime.combine(values['start'], time.min), tz)
    end = timezone.make_aware(datetime.combine(values['end'] + timedelta(days=1), time.min), tz)
    records = scope.filter(paid_at__gte=start, paid_at__lt=end)
    if values['department']:
        records = records.filter(department=values['department'])
    if values['employees']:
        records = records.filter(owner_id__in=set(values['employees']))
    documents = list(records.order_by('-paid_at', '-pk'))
    teams = defaultdict(lambda: {'count': 0, 'total': Decimal('0')})
    people = {}
    total = Decimal('0')
    for document in documents:
        amount = document.total
        total += amount
        document.summary_date = timezone.localtime(document.paid_at, tz).date()
        team = teams[document.department]
        team['count'] += 1
        team['total'] += amount
        person = people.setdefault(document.owner_id, {
            'name': document.owner.first_name or document.owner.username, 'departments': set(), 'count': 0, 'total': Decimal('0'),
        })
        person['departments'].add(document.department)
        person['count'] += 1
        person['total'] += amount
    team_rows = [dict(name=name, **value) for name, value in sorted(teams.items())]
    person_rows = sorted(people.values(), key=lambda row: row['name'])
    for person in person_rows:
        person['department_label'] = ' · '.join(sorted(person['departments']))
    return dict(documents=documents, total=total, count=len(documents), teams=team_rows, people=person_rows)


def workbook(result, values, selected_names):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    book = Workbook()
    summary = book.active
    summary.title = '집계 요약'

    def row(sheet, values, text_columns=(), header=False):
        sheet.append(values)
        for column in text_columns:
            sheet.cell(sheet.max_row, column).data_type = 's'
        for cell in sheet[sheet.max_row]:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if header:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='0B4F3D')

    row(summary, ['집계 기준', '결제 완료 등록일'], (1, 2))
    row(summary, ['기간', f'{values["start"]:%Y-%m-%d} ~ {values["end"]:%Y-%m-%d}'], (1, 2))
    row(summary, ['팀', values['department'] or '전체 팀'], (1, 2))
    row(summary, ['직원', ', '.join(selected_names) or '전체 직원'], (1, 2))
    row(summary, ['총 구매금액', result['total']], (1,))
    row(summary, ['구매내역 건수', result['count']], (1,))
    row(summary, [])
    row(summary, ['팀별 합계', '구매내역', '금액'], (1, 2, 3), header=True)
    for team in result['teams']:
        row(summary, [team['name'], team['count'], team['total']], (1,))
    row(summary, [])
    row(summary, ['직원별 합계', '팀', '구매내역', '금액'], (1, 2, 3, 4), header=True)
    for person in result['people']:
        row(summary, [person['name'], person['department_label'], person['count'], person['total']], (1, 2))
    for column in range(1, 5):
        summary.column_dimensions[get_column_letter(column)].width = 27
    summary.freeze_panes = 'A8'

    detail = book.create_sheet('구매내역')
    row(detail, ['결제완료일', '팀', '기안자', '품목', '구매처', '수량', '단가', '금액', '용도'], tuple(range(1, 10)), header=True)
    for document in result['documents']:
        row(detail, [document.summary_date, document.department, document.owner.first_name or document.owner.username,
            document.product, document.vendor, document.quantity, document.unit_price, document.total, document.reason], (2, 3, 4, 5, 9))
        detail.cell(detail.max_row, 1).number_format = 'yyyy-mm-dd'
        for column in [7, 8]:
            detail.cell(detail.max_row, column).number_format = '#,##0.##'
    detail.freeze_panes = 'A2'
    detail.auto_filter.ref = detail.dimensions
    for column in range(1, 10):
        detail.column_dimensions[get_column_letter(column)].width = 30 if column in [4, 5, 9] else 18
    return book
