"""New purchase indicators are independent of approval and payment completion."""
import unittest
from io import BytesIO
from django.conf import settings
if not settings.INTEGRATED_PURCHASES:raise unittest.SkipTest('Integrated purchase requests only.')
from django.test import TestCase
from django.urls import reverse
from . import test_integrated as fixtures
from .models import PurchaseNotice,PurchaseReport
from .services import submit_report
from uuid import uuid4


class NewRequestIndicatorsTests(TestCase):
    setUpTestData=classmethod(fixtures.IntegratedTests.setUpTestData.__func__)
    user=fixtures.IntegratedTests.user
    as_user=fixtures.IntegratedTests.as_user
    create=fixtures.IntegratedTests.create
    confirmed=fixtures.IntegratedTests.confirmed

    def report(self,docs):return submit_report(self.user('finance'),[d.pk for d in docs],'새 구매요청 시험','업무용 장비',uuid4())
    def status(self):return self.client.get('/notifications/status/').json()

    def test_home_menu_and_new_list_count_reports_not_items_or_approval_tasks(self):
        first=self.report([self.confirmed(),self.confirmed('dev_manager')]);second=self.report([self.confirmed()])
        self.as_user('ceo');home=self.client.get('/')
        self.assertEqual(home.context['counts']['new_purchase'],2);self.assertContains(home,'새 구매요청')
        self.assertContains(home,'data-live-purchase-region="new_purchase"')
        self.assertEqual(self.status()['purchase_counts']['new_purchase'],2);self.assertEqual(self.status()['waiting'],0)
        url=reverse('purchasing:reports')+'?new=1';response=self.client.get(url)
        self.assertEqual(len(response.context['reports']),2)
        self.assertTrue(all(r.is_new_payment for r in response.context['reports']))
        for r in [first,second]:r.refresh_from_db();self.assertIsNone(r.representative_viewed_at)

    def test_only_representative_opening_detail_clears_new_and_preserves_payment_state(self):
        doc=self.confirmed();first=self.report([doc]);second=self.report([self.confirmed()]);url=reverse('purchasing:report_detail',args=[first.pk])
        self.as_user('finance');self.client.get(url);first.refresh_from_db();self.assertIsNone(first.representative_viewed_at)
        self.as_user('quality_manager');self.assertEqual(self.client.get(url).status_code,404)
        first.refresh_from_db();self.assertIsNone(first.representative_viewed_at)
        PurchaseNotice.objects.filter(user=self.user('ceo')).update(read=True)
        self.as_user('ceo');self.assertEqual(self.status()['purchase_counts']['new_purchase'],2)
        self.client.get(url);first.refresh_from_db();doc.refresh_from_db();seen=first.representative_viewed_at
        self.assertIsNotNone(seen);self.assertEqual(first.revision,0);self.assertEqual(first.status,'pending')
        self.assertEqual(first.items.get().decision,'pending');self.assertEqual(doc.status,'ceo');self.assertFalse(doc.completion_unread)
        self.assertEqual(self.status()['purchase_counts']['new_purchase'],1);self.assertEqual(self.status()['waiting'],0)
        self.client.get(url);first.refresh_from_db();self.assertEqual(first.representative_viewed_at,seen)
        response=self.client.get(reverse('purchasing:reports')+'?new=1')
        self.assertEqual([r.pk for r in response.context['reports']],[second.pk])
        self.assertEqual(len(self.client.get(reverse('purchasing:reports')).context['reports']),2)

    def test_empty_indicator_still_has_live_region_and_non_representatives_get_zero(self):
        self.as_user('ceo');response=self.client.get('/')
        self.assertEqual(response.context['counts']['new_purchase'],0)
        self.assertContains(response,'data-live-purchase-region="new_purchase" hidden')
        self.report([self.confirmed()]);self.as_user('finance')
        self.assertEqual(self.status()['purchase_counts']['new_purchase'],0)

    def test_historical_reports_do_not_appear_as_new_purchase_requests(self):
        report=self.report([self.confirmed()]);PurchaseReport.objects.filter(pk=report.pk).update(payment_tracking=False)
        self.as_user('ceo');self.assertEqual(self.status()['purchase_counts']['new_purchase'],0)
        self.assertEqual(self.status()['waiting'],1)

    def test_excel_for_finance_and_representative_removes_three_columns_and_keeps_purpose(self):
        from openpyxl import load_workbook
        doc=self.confirmed();report=self.report([doc]);url=reverse('purchasing:report_export',args=[report.pk])
        for who in ['finance','ceo']:
            self.as_user(who);response=self.client.post(url);self.assertEqual(response.status_code,200)
            sheet=load_workbook(BytesIO(response.content)).active
            self.assertEqual(sheet.max_column,10);self.assertEqual(sheet.cell(1,10).value,'구매사유')
            self.assertEqual(sheet.cell(2,10).value,doc.reason)
            self.assertEqual(sheet.cell(2,9).value,float(doc.total))
            for label in ['진행상황','사유','수령 상태']:self.assertNotIn(label,[cell.value for cell in sheet[1]])
        report.refresh_from_db();self.assertIsNone(report.representative_viewed_at)
