"""Actual payment registration replaces representative approval for new reports."""
import unittest
from io import BytesIO
from uuid import uuid4
from django.conf import settings
if not settings.INTEGRATED_PURCHASES:raise unittest.SkipTest('Integrated purchase payments only.')
from django.test import TestCase, Client
from django.urls import reverse
from django.core.exceptions import PermissionDenied, ValidationError
from workflow.models import Profile
from . import test_integrated as fixtures
from .models import PurchasePolicy, PurchaseRequest, PurchaseReport, PurchaseNotice, PurchaseEvent
from .services import submit_report, decide_report, request_action, save_request, amend_amount
from .payments import process_payment
from .forms import RequestForm


class PaymentWorkflowTests(TestCase):
    setUpTestData=classmethod(fixtures.IntegratedTests.setUpTestData.__func__)
    user=fixtures.IntegratedTests.user
    as_user=fixtures.IntegratedTests.as_user
    create=fixtures.IntegratedTests.create
    confirmed=fixtures.IntegratedTests.confirmed

    def report(self,docs):
        return submit_report(self.user('finance'),[d.pk for d in docs],'시험 결제 요청','업무용 장비',uuid4())

    def process(self,report,items,action='paid',reason='',apply=True,actor=None,revision=None):
        return process_payment(report.pk,actor or self.user('finance'),[i.pk for i in items],action,reason,report.revision if revision is None else revision,apply=apply)

    def test_new_report_is_read_only_for_representative_and_not_an_approval_task(self):
        Profile.objects.filter(user=self.user('ceo')).update(representative_can_approve=False)
        d=self.confirmed();r=self.report([d]);self.assertTrue(r.payment_tracking)
        self.as_user('ceo');url=reverse('purchasing:report_detail',args=[r.pk])
        response=self.client.get(url)
        for label in ['구매 진행현황','기안자','구매처','수량','가격','용도','결제 대기']:self.assertContains(response,label)
        for text in ['data-payment-form','name="items"','선택 품목 결제 완료','수령 현황','대표 결재 대기']:self.assertNotContains(response,text)
        self.assertEqual(self.client.post(url,dict(action='paid',items=[r.items.get().pk],revision=0,intent='confirm')).status_code,403)
        self.assertEqual(self.client.get('/notifications/status/').json()['waiting'],0)
        self.assertNotContains(self.client.get('/documents/?mode=pending'),r.title)
        self.assertContains(self.client.get(reverse('purchasing:reports')),r.title)
        self.assertNotContains(self.client.get(reverse('purchasing:detail',args=[d.pk])),'자료 추가')
        with self.assertRaises(PermissionDenied):decide_report(r.pk,self.user('ceo'),{str(r.items.get().pk):dict(decision='allow')},0,final=True)
        with self.assertRaises(PermissionDenied):self.process(r,[r.items.get()],actor=self.user('ceo'))
        d.refresh_from_db();self.assertEqual(d.status,'ceo');self.assertFalse(d.completion_unread)

    def test_payment_preview_has_no_side_effect_and_only_selected_item_completes(self):
        first=self.confirmed();second=self.confirmed('dev_manager');r=self.report([first,second]);items=list(r.items.all());selected=items[0]
        self.as_user('finance');url=reverse('purchasing:report_detail',args=[r.pk]);before=PurchaseEvent.objects.count()
        self.assertEqual(self.client.get(reverse('purchasing:reports')).context['purchase_tab'],'sent')
        response=self.client.post(url,dict(action='paid',items=[selected.pk],revision=0))
        self.assertContains(response,'선택 품목 합계');self.assertContains(response,'결제 완료 확정')
        self.assertEqual(PurchaseEvent.objects.count(),before);first.refresh_from_db();self.assertEqual(first.status,'ceo')
        self.client.post(url,dict(action='paid',items=[selected.pk],revision=0,intent='confirm'))
        first.refresh_from_db();second.refresh_from_db();r.refresh_from_db()
        self.assertEqual((first.status,second.status,r.status),('paid','ceo','pending'))
        self.assertIsNotNone(first.paid_at);self.assertTrue(first.completion_unread)
        selected.refresh_from_db();self.assertEqual(selected.processed_by,self.user('finance'));self.assertIsNotNone(selected.processed_at)
        notices=list(PurchaseNotice.objects.filter(user=first.owner,request=first,text__contains='최종결재가 완료'))
        self.assertEqual(len(notices),1);self.assertEqual(notices[0].text,f'{first.product} · 최종결재가 완료되었습니다.')
        self.assertFalse(PurchaseNotice.objects.filter(user=first.owner,request=first,text__contains='결제 완료').exists())
        self.assertEqual(r.result_label,'일부 결제 · 결제 대기')
        self.as_user('ceo');self.assertNotContains(self.client.get('/documents/?mode=history'),r.title)
        self.assertContains(self.client.get(url),'결제 완료')

    def test_finance_rejects_one_item_and_requires_reason(self):
        docs=[self.confirmed(),self.confirmed('dev_manager')];r=self.report(docs);items=list(r.items.all())
        with self.assertRaises(ValidationError):self.process(r,[items[1]],'deny')
        self.process(r,[items[1]],'deny','구매처 확인 필요');r.refresh_from_db()
        docs[1].refresh_from_db();self.assertEqual(docs[1].status,'denied');self.assertEqual(docs[1].decision_reason,'구매처 확인 필요')
        self.process(r,[items[0]]);r.refresh_from_db();self.assertEqual(r.status,'decided');self.assertEqual(r.result_label,'처리 완료 · 일부 결제')
        self.assertTrue(PurchaseNotice.objects.filter(user=docs[1].owner,request=docs[1],text__contains='반려').exists())

    def test_finance_can_reject_before_batch_and_cannot_use_that_action_after_send(self):
        d=self.confirmed();self.assertEqual(request_action(d.pk,self.user('finance'),'finance_reject','중복 구매').status,'denied')
        d=self.confirmed();self.report([d])
        with self.assertRaises(ValidationError):request_action(d.pk,self.user('finance'),'finance_reject','화면 우회')

    def test_payment_cannot_be_repeated_or_changed_to_rejection(self):
        d=self.confirmed();r=self.report([d]);item=r.items.get();self.process(r,[item]);r.refresh_from_db()
        for action in ['paid','deny']:
            with self.assertRaises(ValidationError):self.process(r,[item],action,'잘못된 재처리')
        self.assertEqual(PurchaseEvent.objects.filter(request=d,event='결제 완료 확인').count(),1)
        with self.assertRaises(ValidationError):request_action(d.pk,self.user('ceo'),'paid','우회')

    def test_revoked_finance_and_other_finance_cannot_process_original_report(self):
        d=self.confirmed();r=self.report([d]);p=PurchasePolicy.objects.get(pk=1);p.finance.add(self.user('quality_lead'))
        with self.assertRaises(PermissionDenied):self.process(r,[r.items.get()],actor=self.user('quality_lead'))
        p.finance.remove(self.user('finance'))
        with self.assertRaises(PermissionDenied):self.process(r,[r.items.get()])

    def test_foreign_duplicate_and_stale_items_are_rejected_atomically(self):
        d=self.confirmed();r=self.report([d]);other=self.report([self.confirmed('dev_manager')]);item=r.items.get()
        for ids in [[item.pk,other.items.get().pk],[item.pk,item.pk],[]]:
            with self.assertRaises(ValidationError):process_payment(r.pk,self.user('finance'),ids,'paid','',0,apply=True)
        self.process(r,[item]);r.refresh_from_db()
        with self.assertRaises(ValidationError):self.process(r,[item],'undo','등록 오류',revision=0)
        self.assertFalse(PurchaseEvent.objects.filter(report=other,event='결제 완료 확인').exists())

    def test_correction_retains_audit_and_cannot_undo_received_purchase(self):
        d=self.confirmed();r=self.report([d]);item=r.items.get();self.process(r,[item]);r.refresh_from_db()
        with self.assertRaises(ValidationError):self.process(r,[item],'undo')
        self.process(r,[item],'undo','대표님 결제 확인 착오');r.refresh_from_db();d.refresh_from_db()
        self.assertEqual((d.status,r.status),('ceo','pending'));self.assertFalse(d.completion_unread);self.assertIsNone(d.paid_at)
        self.assertTrue(PurchaseEvent.objects.filter(request=d,event='결제 완료 확인').exists())
        self.assertTrue(PurchaseEvent.objects.filter(request=d,event='결제 완료 등록 취소').exists())
        self.process(r,[item]);r.refresh_from_db();request_action(d.pk,d.owner,'receive',receipt_key=uuid4())
        with self.assertRaises(ValidationError):self.process(r,[item],'undo','수령 후 취소 시도')
        self.as_user('finance');self.assertEqual(self.client.get(reverse('purchasing:reports')+'?tab=completed').context['total'],1)

    def test_hold_resubmission_preserves_original_snapshot(self):
        d=self.confirmed();r=self.report([d]);item=r.items.get();snapshot=dict(item.snapshot)
        self.process(r,[item],'hold','규격 확인');request_action(d.pk,self.user('finance'),'prepare_hold','규격 확인 완료')
        d.refresh_from_db();next_report=self.report([d]);self.assertTrue(next_report.payment_tracking)
        item.refresh_from_db();self.assertEqual(item.snapshot,snapshot);self.assertEqual(item.decision,'hold')

    def test_amount_changes_preserve_old_request_and_require_new_payment_request(self):
        for quantity,price in [(3,'60000'),(1,'40000')]:
            d=self.confirmed();r=self.report([d]);item=r.items.get();before=dict(item.snapshot)
            d.refresh_from_db();d=amend_amount(d.pk,self.user('finance'),quantity,price,'구매처 견적 변경',d.revision)
            self.assertEqual(d.status,'reapproval');self.assertFalse(d.completion_unread);r.refresh_from_db();item.refresh_from_db()
            self.assertEqual(item.snapshot,before);self.assertEqual(item.decision,'hold')
            with self.assertRaises(ValidationError):self.process(r,[item])
            next_report=self.report([d]);next_item=next_report.items.get()
            self.assertEqual(next_item.snapshot['quantity'],quantity);self.assertIn('amendment',next_item.snapshot)
            self.process(next_report,[next_item]);d.refresh_from_db();self.assertEqual(d.status,'paid')
            with self.assertRaises(ValidationError):amend_amount(d.pk,self.user('finance'),1,'1','결제 후 변경 시도',d.revision)

    def test_csrf_and_owner_access_are_enforced(self):
        d=self.confirmed();r=self.report([d]);url=reverse('purchasing:report_detail',args=[r.pk])
        self.as_user('quality_manager');self.assertEqual(self.client.get(url).status_code,404)
        c=Client(enforce_csrf_checks=True);c.force_login(self.user('finance'))
        self.assertEqual(c.post(url,dict(action='paid',items=[r.items.get().pk],revision=0,intent='confirm')).status_code,403)

    def test_reviewer_picker_visible_and_completion_secondary_only_when_authorized(self):
        d=self.create('quality','review');url=reverse('purchasing:detail',args=[d.pk]);self.as_user('quality_manager')
        response=self.client.get(url);self.assertContains(response,'data-forward-visible');self.assertNotContains(response,'<details class="review-forward"')
        self.assertNotContains(response,'value="review_final"');self.assertContains(response,'<section class="secondary-actions"')
        PurchasePolicy.objects.get(pk=1).review_final_users.add(self.user('quality_manager'))
        response=self.client.get(url);self.assertContains(response,'<button class="full">전결로 승인완료</button>');self.assertContains(response,'<button class="primary full">선택한 승인자에게 전달</button>')
        request_action(d.pk,self.user('quality_manager'),'review_final');d.refresh_from_db();self.assertEqual(d.status,'ready');self.assertFalse(d.completion_unread)

    def test_link_required_to_submit_but_drafts_and_old_blank_links_preserved(self):
        d=self.create();values=dict(product=d.product,vendor=d.vendor,quantity=d.quantity,unit_price=d.unit_price,reason=d.reason,request_method='direct',url='')
        with self.assertRaises(ValidationError):save_request(self.user('quality_manager'),values,submit=True,key=uuid4())
        draft=save_request(self.user('quality_manager'),values,key=uuid4());self.assertEqual(draft.status,'draft')
        data={**values,'key':str(uuid4()),'intent':'submit'}
        f=RequestForm(data,user=self.user('quality_manager'),policy=PurchasePolicy.objects.get(pk=1));self.assertFalse(f.is_valid());self.assertIn('url',f.errors)
        data['intent']='draft';f=RequestForm(data,user=self.user('quality_manager'),policy=PurchasePolicy.objects.get(pk=1));self.assertTrue(f.is_valid(),f.errors)
        PurchaseRequest.objects.filter(pk=d.pk).update(url='');d.refresh_from_db();self.assertEqual(d.url,'')

    def test_representative_export_contains_only_purchase_details(self):
        from openpyxl import load_workbook
        d=self.confirmed();r=self.report([d]);self.as_user('ceo')
        response=self.client.post(reverse('purchasing:report_export',args=[r.pk]));self.assertEqual(response.status_code,200)
        sheet=load_workbook(BytesIO(response.content)).active;header=[c.value for c in sheet[1]]
        self.assertEqual(sheet.max_column,10);self.assertEqual(header[-1],'구매사유')
        for label in ['진행상황','사유','수령 상태']:self.assertNotIn(label,header)
