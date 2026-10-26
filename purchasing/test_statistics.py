"""Accounting totals preserve dates, scope, decimal precision and export filters."""
import unittest
from datetime import datetime, date, time, timedelta
from decimal import Decimal
from io import BytesIO
from urllib.parse import parse_qs, urlsplit
from uuid import uuid4

from django.conf import settings
if not settings.INTEGRATED_PURCHASES:raise unittest.SkipTest('Integrated purchase summary only.')
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from . import test_integrated as fixtures
from .models import PurchaseRequest, PurchasePolicy, PurchaseEvent
from .services import request_action, submit_report
from .payments import process_payment


@override_settings(TIME_ZONE='Asia/Seoul')
class PurchaseSummaryTests(TestCase):
    setUpTestData=classmethod(fixtures.IntegratedTests.setUpTestData.__func__)
    user=fixtures.IntegratedTests.user
    as_user=fixtures.IntegratedTests.as_user
    create=fixtures.IntegratedTests.create

    def paid(self,who='quality_manager',day='2026-10-10',quantity=1,unit_price='100.00',**overrides):
        document=self.create(who,'direct',quantity=quantity,unit_price=Decimal(unit_price),**overrides)
        if not document.confirmed_at:document=request_action(document.pk,document.purchase_checker,'confirm_items')
        report=submit_report(self.user('finance'),[document.pk],'집계용 가상 요청','업무용',uuid4())
        process_payment(report.pk,self.user('finance'),[report.items.get().pk],'paid','',0,apply=True)
        stamp=timezone.make_aware(datetime.combine(date.fromisoformat(day),time(12)),timezone.get_current_timezone())
        PurchaseRequest.objects.filter(pk=document.pk).update(paid_at=stamp)
        document.refresh_from_db()
        return document

    def query(self,**changes):
        values=dict(start='2026-10-01',end='2026-10-31',department='')
        values.update(changes)
        return values

    def summary(self,**changes):
        self.as_user('finance')
        return self.client.get(reverse('purchasing:summary'),self.query(**changes))

    def test_summary_and_download_are_finance_only_and_permission_is_live(self):
        self.paid()
        for who in ['quality','ceo','admin']:
            self.as_user(who)
            for download in ['', '1']:
                self.assertEqual(self.client.get(reverse('purchasing:summary'),self.query(download=download)).status_code,403)
        self.as_user('finance');self.assertEqual(self.client.get(reverse('purchasing:summary')).status_code,200)
        PurchasePolicy.objects.get(pk=1).finance.remove(self.user('finance'))
        self.assertEqual(self.client.get(reverse('purchasing:summary')).status_code,403)

    def test_period_is_inclusive_in_local_timezone_and_uses_payment_registration_date(self):
        midnight=self.paid(day='2026-09-30');late=self.paid();before=self.paid();after=self.paid()
        tz=timezone.get_current_timezone()
        moments=[datetime(2026,10,1),datetime(2026,10,1,23,59,59,999999),datetime(2026,9,30,23,59,59,999999),datetime(2026,10,2)]
        for document,stamp in zip([midnight,late,before,after],moments):PurchaseRequest.objects.filter(pk=document.pk).update(paid_at=timezone.make_aware(stamp,tz))
        response=self.summary(start='2026-10-01',end='2026-10-01')
        self.assertEqual({d.pk for d in response.context['result']['documents']},{midnight.pk,late.pk})
        self.assertEqual(response.context['result']['total'],Decimal('200.00'))

    def test_team_and_multiple_employee_filters_keep_exact_totals_without_duplicates(self):
        first=self.paid('quality_manager',quantity=2,unit_price='0.10')
        second=self.paid('quality_lead',quantity=1,unit_price='0.20')
        third=self.paid('dev_manager',quantity=3,unit_price='100')
        self.assertEqual(self.summary().context['result']['total'],Decimal('300.40'))
        result=self.summary(department='품질팀',employees=[first.owner_id,second.owner_id,first.owner_id]).context['result']
        self.assertEqual((result['count'],result['total']),(2,Decimal('0.40')))
        self.assertEqual(len(result['people']),2);self.assertEqual(len(result['teams']),1)
        result=self.summary(employees=[second.owner_id,third.owner_id]).context['result']
        self.assertEqual(result['total'],Decimal('300.20'));self.assertEqual(len(result['teams']),2)

    def test_pending_rejected_held_and_unknown_payment_dates_are_excluded(self):
        paid=self.paid();received=self.paid();request_action(received.pk,received.owner,'receive',receipt_key=uuid4())
        for state in ['ceo','hold','denied','rejected','cancelled','ready']:
            document=self.paid();PurchaseRequest.objects.filter(pk=document.pk).update(status=state)
        unknown=self.paid();PurchaseRequest.objects.filter(pk=unknown.pk).update(status='received',paid_at=None)
        result=self.summary().context['result']
        self.assertEqual({d.pk for d in result['documents']},{paid.pk,received.pk})

    def test_invalid_period_team_and_employee_never_fall_back_to_all_records(self):
        self.paid();foreign=self.user('quality').pk
        for changes in [dict(start='bad'),dict(start='2026-11-01'),dict(end=''),dict(end='9999-12-31'),dict(department='없는팀'),dict(employees=[foreign])]:
            response=self.summary(**changes)
            self.assertFalse(response.context['filters_valid']);self.assertEqual(response.context['result']['count'],0)
            self.assertFalse(response.context['export_url'])
            response=self.summary(download='1',**changes);self.assertEqual(response.status_code,400)

    def test_employee_must_belong_to_selected_historical_department(self):
        document=self.paid('dev_manager')
        document.owner.profile.department='품질팀';document.owner.profile.save(update_fields=['department'])
        response=self.summary(department='개발팀',employees=[document.owner_id])
        self.assertTrue(response.context['filters_valid']);self.assertEqual(response.context['result']['count'],1)
        self.paid('quality_manager')
        self.assertFalse(self.summary(department='품질팀',employees=[document.owner_id]).context['filters_valid'])

    def test_other_accountants_scope_is_not_leaked_through_choices_or_download(self):
        visible=self.paid();hidden=self.paid('dev_manager',product='권한 밖 품목')
        other=self.user('quality_lead');PurchasePolicy.objects.get(pk=1).finance.add(other)
        route=dict(hidden.purchase_route);route['collector_ids']=[other.pk]
        PurchaseRequest.objects.filter(pk=hidden.pk).update(purchase_route=route)
        hidden.current_report.created_by=other;hidden.current_report.save(update_fields=['created_by'])
        response=self.summary()
        self.assertEqual([d.pk for d in response.context['result']['documents']],[visible.pk])
        self.assertNotContains(response,'권한 밖 품목')
        self.assertNotIn(str(hidden.owner_id),[person['id'] for person in response.context['form'].people])
        self.assertFalse(self.summary(employees=[hidden.owner_id]).context['filters_valid'])

    def test_pagination_totals_and_excel_cover_all_filtered_records(self):
        for _ in range(26):self.paid(quantity=2,unit_price='100')
        self.paid('dev_manager',unit_price='999')
        response=self.summary(department='품질팀',employees=[self.user('quality_manager').pk],page=2)
        self.assertEqual((response.context['result']['count'],response.context['result']['total']),(26,Decimal('5200.00')))
        self.assertEqual(len(response.context['page_obj']),1)
        params=parse_qs(urlsplit(response.context['export_url']).query)
        self.assertEqual(params['employees'],[str(self.user('quality_manager').pk)])
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(self.client.get(response.context['export_url']).content))
        self.assertEqual(book['구매내역'].max_row,27);self.assertEqual(book['집계 요약']['B5'].value,5200)

    def test_workbook_contains_filtered_summary_and_details_and_escapes_text_formulas(self):
        from openpyxl import load_workbook
        first=self.paid(product='=1+1',vendor='=SUM(1,1)',reason='업무용 테스트',quantity=3,unit_price='123.45')
        self.paid('dev_manager',unit_price='999')
        response=self.summary(employees=[first.owner_id],download='1')
        self.assertEqual(response.status_code,200)
        book=load_workbook(BytesIO(response.content));self.assertEqual(book.sheetnames,['집계 요약','구매내역'])
        sheet=book['구매내역'];self.assertEqual(sheet.max_column,9);self.assertEqual(sheet.max_row,2)
        self.assertEqual(sheet['D2'].value,'=1+1');self.assertEqual(sheet['D2'].data_type,'s')
        self.assertEqual(sheet['E2'].data_type,'s');self.assertEqual(sheet['H2'].value,370.35);self.assertEqual(sheet['I2'].value,'업무용 테스트')
        for label in ['진행상황','처리 사유','수령 상태']:self.assertNotIn(label,[cell.value for cell in sheet[1]])

    def test_reading_summary_does_not_change_payment_or_receipt_state(self):
        document=self.paid();before=(document.status,document.paid_at,document.received_at,document.revision)
        events=PurchaseEvent.objects.count();self.summary();document.refresh_from_db()
        self.assertEqual((document.status,document.paid_at,document.received_at,document.revision),before)
        self.assertEqual(PurchaseEvent.objects.count(),events)

    def test_default_period_is_this_month(self):
        self.as_user('finance');response=self.client.get(reverse('purchasing:summary'));today=timezone.localdate()
        self.assertEqual(response.context['form'].cleaned_data['start'],today.replace(day=1))
        self.assertEqual(response.context['form'].cleaned_data['end'],today)
