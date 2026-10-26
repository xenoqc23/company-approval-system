from uuid import uuid4
from django.conf import settings
from django.urls import reverse
from decimal import Decimal
from django.contrib import messages
from django.contrib.auth import authenticate,login,logout,get_user_model
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied,ValidationError
from django.db import transaction,IntegrityError
from django.db.models import Q,Sum,F,DecimalField,ExpressionWrapper
from django.http import JsonResponse,HttpResponse,FileResponse,Http404
from django.shortcuts import render,redirect,get_object_or_404
from django.utils import timezone
from django.utils.cache import add_never_cache_headers
from django.views.decorators.http import require_POST,require_GET
from django.views.decorators.debug import sensitive_post_parameters
from workflow.models import RANKS
from .models import PurchaseRequest,PurchaseReport,PurchaseReportItem,PurchaseNotice,PurchasePolicy,DemoAccount,PurchasePermission
from .forms import RequestForm,PolicyForm,AttachmentForm,SupplementForm,AmountChangeForm,PurchaseForwardForm
from .access import active,is_admin,is_finance,requests_for,reports_for,require_finance,policy,notices_for,is_confirmer,can_view_purchase,finance_assigned,finance_requests,checker_assigned,checker_requests
from .services import save_request,request_action,submit_report,decide_report,event,supplement_request,amend_amount,reviewer_can_finish,combines_confirmation
from .attachments import StoredFiles,visible_attachments,can_upload
from .context import demo_choices
from .fulfillment import receipt_info,report_fulfillment,report_receipt_label
from .progress import progress, PURCHASING_STATES, REQUESTER_STATES, visible_notice, events_for, approval_track, purchase_progress, pending_columns, DISPLAY_LABELS, task_for


def page(request,template,title,**context):
    section='purchase_check' if context.get('mode')=='completed' else 'regulations' if template=='regulations' else 'staff' if template=='staff' else 'compose' if template=='compose' else 'purchase_inbox'
    response=render(request,f'purchasing/{template}.html',dict(title=title,page_name=template,section=section,**context))
    add_never_cache_headers(response);return response


@sensitive_post_parameters('password')
def login_view(request):
    if settings.INTEGRATED_PURCHASES:return redirect('dashboard' if request.user.is_authenticated else 'login')
    if request.user.is_authenticated:return redirect('purchasing:home')
    if request.method=='POST':
        u=authenticate(request,username=request.POST.get('username',''),password=request.POST.get('password',''))
        if u and active(u):login(request,u);return redirect('purchasing:home')
        messages.error(request,'아이디와 비밀번호를 확인해 주세요.')
    return page(request,'login','구매요청 로컬 테스트 로그인',accounts=demo_choices())


@login_required
@require_POST
def logout_view(request):logout(request);return redirect('login' if settings.INTEGRATED_PURCHASES else 'purchasing:login')


@login_required
@require_POST
def switch(request):
    # Only the fictional accounts explicitly created by the isolated demo seeder.
    from django.conf import settings
    if not settings.PURCHASE_LAB_ENABLED or settings.DATA_DIR.name!='purchase-test':raise PermissionDenied
    from .context import DEMO_ROLES
    account=get_object_or_404(DemoAccount.objects.filter(user__username__in=[a.user.username for a in demo_choices()]),user_id=request.POST.get('user') if str(request.POST.get('user','')).isdigit() else None)
    if not active(account.user):raise PermissionDenied
    login(request,account.user,backend='django.contrib.auth.backends.ModelBackend')
    return redirect('dashboard' if settings.INTEGRATED_PURCHASES else 'purchasing:home')


@login_required
def home(request):
    if settings.INTEGRATED_PURCHASES:return redirect('dashboard')
    qs=PurchaseRequest.objects.filter(owner=request.user)
    recent=list(qs.select_related('reviewer','approver','director')[:8])
    for d in recent:d.progress_info=progress(d, simplified=True)
    stats=dict(internal=qs.filter(status__in=['review','approval','director']).count(),purchasing=qs.filter(status__in=PURCHASING_STATES).count(),active=qs.filter(status__in=['review','approval','director',*PURCHASING_STATES]).count(),
               allowed=qs.filter(status__in=['allowed','paid']).count(),denied=qs.filter(status__in=['denied','hold','rejected']).count(),
               own_count=qs.count(),received=qs.filter(status='received').count())
    notices=list(notices_for(request.user).select_related('request','report')[:6])
    for n in notices:n.display_text=visible_notice(n,request.user)
    return page(request,'home','내 업무',recent=recent,stats=stats,notices=notices)


@login_required
def request_list(request, purchase_tab=None):
    mode=('hold' if purchase_tab=='hold' else 'finance') if purchase_tab else request.GET.get('mode','mine');u=request.user
    qs=requests_for(u).select_related('owner','reviewer','approver','director')
    titles=dict(mine='내 구매요청',inbox='내가 처리할 구매요청',director='이사 최종승인 목록',completed='구매 확인함',finance='구매 모아보기',hold='구매보류함',all='구매요청 전체')
    if mode=='mine':qs=qs.filter(owner=u)
    elif mode=='inbox':qs=qs.filter(Q(status='review',reviewer=u)|Q(status='approval',approver=u)|Q(status='director',director=u))
    elif mode=='director':qs=qs.filter(director=u).exclude(status='draft')
    elif mode=='completed':
        if not is_confirmer(u):raise PermissionDenied
        qs=qs.filter(pk__in=checker_requests(u).values('pk'))
    elif mode=='finance':
        require_finance(u);qs=qs.filter(pk__in=finance_requests(u).values('pk'))
    elif mode=='hold':
        require_finance(u);qs=qs.filter(pk__in=finance_requests(u).values('pk'),status='hold')
    elif mode=='all':
        if not is_admin(u):raise PermissionDenied
    else:raise PermissionDenied
    queue_counts=dict(approved=qs.count(),confirm=qs.filter(status='ready',confirmed_at=None).count(),collect=qs.filter(status__in=['ready','reapproval'],confirmed_at__isnull=False).count()) if mode in ['completed','finance'] else {}
    stage=request.GET.get('stage','') if mode in ['completed','finance'] else ''
    if stage=='confirm':qs=qs.filter(status='ready',confirmed_at=None)
    elif stage=='collect':qs=qs.filter(Q(status__in=['ready','reapproval'])|Q(status='hold',hold_ready=True)).filter(confirmed_at__isnull=False)
    receipt=request.GET.get('receipt','') if mode in ['finance','hold'] else ''
    if receipt=='pending':qs=qs.filter(status__in=['allowed','paid'])
    elif receipt=='partial':qs=qs.filter(status__in=['allowed','paid'],received_quantity__gt=0)
    elif receipt=='done':qs=qs.filter(status='received')
    if settings.INTEGRATED_PURCHASES and purchase_tab=='items':
        qs=qs.filter(Q(status__in=['ready','reapproval'])|Q(status='hold',hold_ready=True)).filter(confirmed_at__isnull=False)
    elif settings.INTEGRATED_PURCHASES and purchase_tab=='completed':qs=qs.filter(status='received')
    selected=request.GET.get('state','')
    states=REQUESTER_STATES if mode=='mine' else list(DISPLAY_LABELS.items())
    if settings.INTEGRATED_PURCHASES:
        from .integration import filter_states
        states,groups=filter_states(mode)
        if selected in groups:qs=qs.filter(**groups[selected])
        elif selected:qs=qs.none()
    elif mode=='mine' and selected=='purchasing':qs=qs.filter(status__in=PURCHASING_STATES)
    elif mode=='mine' and selected=='purchase_approved':qs=qs.filter(status__in=['allowed','paid','received'])
    elif selected in dict(PurchaseRequest.STATES):qs=qs.filter(status=selected)
    if mode=='mine' and request.GET.get('unread')=='1':qs=qs.filter(completion_unread=True)
    q=request.GET.get('q','').strip()[:120]
    if q:qs=qs.filter(Q(product__icontains=q)|Q(vendor__icontains=q)|Q(owner__first_name__icontains=q))
    if mode=='inbox':qs=qs.order_by('-urgent',F('needed_by').asc(nulls_last=True),'-created_at','-pk')
    docs=list(qs.select_related('confirmed_by','purchase_checker').prefetch_related('receipts__received_by')[:100])
    for d in docs:
        d.progress_info=progress(d,simplified=not can_view_purchase(u,d));d.queue_info=pending_columns(d);d.receipt_info=receipt_info(d)
        d.can_collect=bool(d.confirmed_at and (d.status in ['ready','reapproval'] or d.status=='hold' and d.hold_ready))
    return page(request,'requests','구매요청함' if purchase_tab else titles[mode],documents=docs,mode=mode,states=states,purchase_tab=purchase_tab,show_states=not purchase_tab,
                inbox_counts=inbox_counts(u) if settings.INTEGRATED_PURCHASES and is_finance(u) else {},
                selected=selected,total=qs.count(),q=q,queue_counts=queue_counts,stage=stage,receipt=receipt,can_batch=mode in ['finance','hold'] and any(d.can_collect for d in docs))


@login_required
def compose(request,pk=None):
    if request.user.profile.is_representative:raise PermissionDenied
    instance=get_object_or_404(PurchaseRequest,pk=pk,owner=request.user,status='draft') if pk else None
    initial={}
    if settings.INTEGRATED_PURCHASES and request.method=='GET' and not pk and request.GET.get('source'):
        from workflow.models import Document
        source=get_object_or_404(Document,pk=request.GET['source'] if request.GET['source'].isascii() and request.GET['source'].isdigit() and len(request.GET['source'])<=18 else None,owner=request.user,kind='office',status__in=['rejected','cancelled'])
        initial=dict(product=source.product,quantity=source.quantity,unit_price=source.unit_price,url=source.url,reason=source.reason,private=source.purchase_private,needed_by=source.needed_date)
    p=policy();form=RequestForm(request.POST if request.method=='POST' else None,request.FILES if request.method=='POST' else None,instance=instance,user=request.user,policy=p,initial=initial)
    if request.method=='POST' and form.is_valid():
        try:
            # Files and metadata form one request: failed uploads must not leave a submitted document.
            with StoredFiles() as stored,transaction.atomic():
                d=save_request(request.user,form.cleaned_data,submit=request.POST.get('intent')=='submit',pk=pk,key=form.cleaned_data['key'])
                stored.add(d,request.user,form.cleaned_data['files'],'request')
            messages.success(request,'구매요청을 제출했습니다.' if d.status!='draft' else '임시저장했습니다.')
            return redirect('purchasing:detail',pk=d.pk)
        except ValidationError as exc:form.add_error(None,exc)
    director=form.checker
    return page(request,'compose','구매요청 작성',form=form,director=director,purchase_configured=bool(p and p.enabled and p.dev_director_id and p.production_director_id and p.representative_id and p.finance.exists()),existing_files=visible_attachments(request.user,instance) if instance else [])


@login_required
def detail(request,pk):
    d=get_object_or_404(requests_for(request.user).select_related('owner','reviewer','approver','director','current_report'),pk=pk)
    notices_for(request.user).filter(request=d,read=False).update(read=True)
    if d.owner_id==request.user.pk and d.completion_unread:
        PurchaseRequest.objects.filter(pk=d.pk,revision=d.revision).update(completion_unread=False)
    assigned={'review':d.reviewer_id,'approval':d.approver_id,'director':d.director_id}.get(d.status)
    simplified=not can_view_purchase(request.user,d)
    permissions=dict(
        can_confirm=checker_assigned(request.user,d) and request.user.pk!=d.owner_id and d.status in ['ready','hold'] and not d.confirmed_at,
        can_approve=assigned==request.user.pk and not (d.status=='review' and d.route.get('reviewer_decides')),
        can_forward=assigned==request.user.pk and d.status=='review' and bool(d.route.get('reviewer_decides')),
        can_review_final=bool(d.route.get('reviewer_decides')) and reviewer_can_finish(request.user,d),
        can_reject=assigned==request.user.pk,
        can_cancel=d.owner_id==request.user.pk and d.status in ['draft','review','approval','director','ready'] and not d.current_report_id and not d.confirmed_at,
        can_receive=d.owner_id==request.user.pk and d.status in ['allowed','paid'],
        can_supplement=d.status=='supplement' and d.owner_id==request.user.pk,
        can_request_supplement=checker_assigned(request.user,d) and d.owner_id!=request.user.pk and d.status=='ready' and not d.confirmed_at and not d.current_report_id,
        can_prepare_hold=finance_assigned(request.user,d) and d.status=='hold',
        can_amend=finance_assigned(request.user,d) and (d.status=='allowed' or d.status in ['ceo','hold'] and d.current_report_id and d.current_report.payment_tracking and d.current_report.created_by_id==request.user.pk) and not d.received_quantity and not d.paid_at,
        can_collect=finance_assigned(request.user,d) and d.status in ['ready','reapproval'] and bool(d.confirmed_at),
        can_finance_reject=finance_assigned(request.user,d) and d.status in ['ready','reapproval','hold'],
        can_decide_purchase=d.status=='ceo' and request.user.profile.is_representative and d.current_report_id and not d.current_report.payment_tracking and d.current_report.representative_id==request.user.pk,
        can_record_payment=d.status=='ceo' and finance_assigned(request.user,d) and d.current_report_id and d.current_report.payment_tracking and d.current_report.created_by_id==request.user.pk)
    has_secondary=any(permissions[key] for key in ['can_cancel','can_reject','can_request_supplement','can_amend','can_finance_reject']) or d.owner_id==request.user.pk and bool(d.received_quantity)
    return page(request,'detail',d.number,doc=d,track=progress(d,simplified=simplified),simplified=simplified,
        visible_events=events_for(d,simplified),approval_cards=approval_track(d),
        purchase_steps=(purchase_progress(d)[:3] if request.user.profile.is_representative else purchase_progress(d)) if not simplified and (d.confirmed_at or d.current_report_id) else [],
        task=task_for(d,request.user,permissions),has_secondary=has_secondary,**permissions,
        forward_form=PurchaseForwardForm(document=d) if permissions['can_forward'] else None,
        approve_combined=permissions['can_approve'] and d.status in ['approval','director'] and combines_confirmation(d,request.user),
        supplements=d.supplements.all(),amendments=d.amendments.all(),receipt_key=uuid4(),
        attachment_form=AttachmentForm(),attachment_list=visible_attachments(request.user,d),can_attach=can_upload(request.user,d),
        can_report=bool(d.current_report_id and reports_for(request.user).filter(pk=d.current_report_id).exists()))


@login_required
@require_POST
def act(request,pk):
    document=get_object_or_404(requests_for(request.user),pk=pk)
    try:
        action=request.POST.get('action','');person=None
        if action=='forward':
            form=PurchaseForwardForm(request.POST,document=document)
            if not form.is_valid():raise ValidationError('승인자 부서와 직원을 확인해 주세요. 기안자·검토자 본인은 선택할 수 없습니다.')
            person=form.cleaned_data['new_person']
        d=request_action(pk,request.user,action,request.POST.get('reason',''),quantity=request.POST.get('quantity'),receipt_key=request.POST.get('receipt_key'),storage_location=request.POST.get('storage_location'),new_person=person)
        text='처리 결과를 저장했습니다.'
        if action=='forward' or action=='approve' and d.status=='approval':text=f'{d.approver.first_name}님에게 승인 요청을 보냈습니다.'
        elif action in ['review_final','approve'] and d.status=='ready':text='승인완료되어 '+('구매 담당자' if d.confirmed_at else '구매확인 담당자')+'에게 전달했습니다.'
        elif action=='confirm_items':text='구매 내용을 확인하고 구매 담당자에게 전달했습니다.'
        elif action=='receive':text='수령 완료로 등록했습니다.'
        elif action=='reject':text='반려 처리하고 기안자에게 알렸습니다.'
        messages.success(request,text)
    except ValidationError as exc:messages.error(request,' '.join(exc.messages))
    return redirect('purchasing:detail',pk=pk)


def selected_ids(raw):
    if not raw or len(raw)>100 or any(not x.isascii() or not x.isdigit() or len(x)>18 or int(x)<1 for x in raw) or len(set(raw))!=len(raw):
        raise ValidationError('중복 없이 1~100개 품목을 선택해 주세요.')
    return [int(x) for x in raw]


@login_required
@require_POST
def report_new(request):
    require_finance(request.user)
    try:
        ids=selected_ids(request.POST.getlist('requests'))
        docs=list(requests_for(request.user).filter(pk__in=ids,status__in=['ready','hold','reapproval'],confirmed_at__isnull=False).select_related('owner').order_by('department','pk'))
        if len(docs)!=len(ids) or any(not finance_assigned(request.user,d) or d.status=='hold' and not d.hold_ready for d in docs):raise ValidationError('이미 요청했거나 현재 요청할 수 없는 품목이 있습니다.')
        if request.POST.get('intent')=='submit':
            report=submit_report(request.user,ids,request.POST.get('title',''),request.POST.get('note',''),request.POST.get('key'))
            messages.success(request,'대표님에게 결제할 구매 내역을 전달했습니다.' if report.payment_tracking else '대표님에게 품목별 구매승인을 요청했습니다.')
            return redirect('purchasing:report_detail',pk=report.pk)
    except ValidationError as exc:
        messages.error(request,' '.join(exc.messages));return redirect('/purchase-lab/requests/?mode=finance')
    return page(request,'report_new','대표님께 결제 요청',documents=docs,total=sum((d.total for d in docs),Decimal('0')),
                key=uuid4(),report_title=f'{timezone.localdate():%Y.%m.%d} 결제 요청')


@login_required
def report_list(request):
    if settings.INTEGRATED_PURCHASES and not (is_finance(request.user) or request.user.profile.is_representative):raise PermissionDenied
    tab=request.GET.get('tab','items') if is_finance(request.user) and settings.INTEGRATED_PURCHASES else 'sent'
    if is_finance(request.user) and settings.INTEGRATED_PURCHASES and 'tab' not in request.GET:
        ready=finance_requests(request.user).filter(Q(status__in=['ready','reapproval'])|Q(status='hold',hold_ready=True)).filter(confirmed_at__isnull=False).exists()
        if not ready and reports_for(request.user).filter(payment_tracking=True,status='pending').exists():tab='sent'
    if is_finance(request.user) and tab in ['items','hold','completed']:return request_list(request,purchase_tab=tab)
    qs=reports_for(request.user).prefetch_related('items__request__receipts__received_by')
    only_new=request.user.profile.is_representative and request.GET.get('new')=='1'
    representative_tab='progress'
    progress_count=completed_count=0
    if request.user.profile.is_representative:
        from .report_status import with_open_items
        qs=with_open_items(qs)
        progress_count=qs.filter(has_open_items=True).count()
        completed_count=qs.filter(has_open_items=False).count()
        representative_tab='completed' if request.GET.get('tab')=='completed' and not only_new else 'progress'
        qs=qs.filter(has_open_items=representative_tab=='progress')
    if only_new:qs=qs.filter(payment_tracking=True,status='pending',representative_viewed_at=None)
    reports=list(qs[:100])
    for report in reports:
        report.is_new_payment=report.payment_tracking and report.status=='pending' and report.representative_viewed_at is None
        items=list(report.items.all());report.fulfillment=report_fulfillment(report,items)
        report.receipt_label=report_receipt_label(report.fulfillment)
        report.requesters=' · '.join(dict.fromkeys(i.snapshot.get('owner','') for i in items))
        report.vendors=' · '.join(dict.fromkeys(i.snapshot.get('vendor','') for i in items))
    return page(request,'reports','구매 진행현황' if request.user.profile.is_representative else '구매요청함',reports=reports,purchase_tab='sent',only_new=only_new,representative_tab=representative_tab,progress_count=progress_count,completed_count=completed_count,inbox_counts=inbox_counts(request.user) if settings.INTEGRATED_PURCHASES and is_finance(request.user) else {})


@login_required
def report_detail(request,pk):
    report=get_object_or_404(reports_for(request.user).select_related('created_by','representative'),pk=pk)
    if request.method=='GET' and report.payment_tracking and request.user.profile.is_representative and report.representative_id==request.user.pk and report.representative_viewed_at is None:
        report.representative_viewed_at=timezone.now()
        PurchaseReport.objects.filter(pk=report.pk,representative=request.user,representative_viewed_at=None).update(representative_viewed_at=report.representative_viewed_at)
    notices_for(request.user).filter(report=report,read=False).update(read=True)
    items=list(report.items.select_related('request').prefetch_related('request__receipts__received_by'))
    for item in items:item.visible_files=visible_attachments(request.user,item.request)
    if report.payment_tracking:
        return payment_detail(request,report,items)
    can_decide=report.status=='pending' and report.representative_id==request.user.pk and request.user.profile.is_representative and (not settings.INTEGRATED_PURCHASES or request.user.profile.representative_can_approve)
    if request.method=='POST':
        if not can_decide:raise PermissionDenied
        try:
            revision=int(request.POST.get('revision','-1'))
            decisions={str(i.pk):dict(decision=request.POST.get(f'decision_{i.pk}','pending'),reason=request.POST.get(f'reason_{i.pk}','')) for i in items}
            decide_report(pk,request.user,decisions,revision,final=request.POST.get('intent')=='final')
            messages.success(request,'품목별 결정을 확정했습니다.' if request.POST.get('intent')=='final' else '검토 내용을 임시저장했습니다.')
            return redirect('purchasing:report_detail',pk=pk)
        except (ValueError,ValidationError) as exc:
            messages.error(request,' '.join(exc.messages) if isinstance(exc,ValidationError) else '화면을 다시 열어 주세요.')
            for item in items:
                choice=request.POST.get(f'decision_{item.pk}','pending')
                if choice in dict(PurchaseReportItem.DECISIONS):item.decision=choice;item.reason=request.POST.get(f'reason_{item.pk}','')[:1000]
    fulfillment=report_fulfillment(report,items)
    return page(request,'report_detail',report.number,fulfillment=fulfillment,receipt_label=report_receipt_label(fulfillment),show_receipts=not request.user.profile.is_representative,report=report,items=items,can_decide=can_decide,
                decisions=[('allow','승인'),('deny','반려'),('hold','보류')],total=sum((i.total for i in items),Decimal('0')),
                allowed_total=sum((i.total for i in items if i.decision=='allow'),Decimal('0')),
                denied_total=sum((i.total for i in items if i.decision=='deny'),Decimal('0')),
                hold_total=sum((i.total for i in items if i.decision=='hold'),Decimal('0')))


def payment_detail(request,report,items):
    from .payments import process_payment
    can_process=is_finance(request.user) and report.created_by_id==request.user.pk
    if request.method=='POST':
        if not can_process:raise PermissionDenied('대표님은 구매 내역을 조회합니다.')
        try:
            summary=process_payment(report.pk,request.user,request.POST.getlist('items'),request.POST.get('action',''),request.POST.get('reason',''),int(request.POST.get('revision','-1')),apply=request.POST.get('intent')=='confirm')
            if request.POST.get('intent')=='confirm':
                messages.success(request,f'{len(summary["items"])}개 품목을 {summary["action_label"]} 처리했습니다.')
                return redirect('purchasing:report_detail',pk=report.pk)
            return page(request,'payment_confirm','선택 품목 확인',**summary)
        except (ValueError,ValidationError) as exc:
            messages.error(request,' '.join(exc.messages) if isinstance(exc,ValidationError) else '화면을 다시 열어 주세요.')
            return redirect('purchasing:report_detail',pk=report.pk)
    for item in items:
        item.request.receipt_info=receipt_info(item.request)
        item.payment_label={'pending':'결제 대기','allow':'결제 완료','deny':'반려','hold':'보류'}[item.decision]
        item.selectable=can_process and item.decision=='pending'
        item.undoable=can_process and item.decision=='allow' and item.request.status=='paid' and not item.request.received_quantity and item.request.current_report_id==report.pk
    from .report_status import with_open_items
    completed=not with_open_items(PurchaseReport.objects.filter(pk=report.pk)).get().has_open_items
    return page(request,'payment_detail',report.title,report=report,items=items,can_process=can_process,
        representative_completed=completed,
        can_undo=any(i.undoable for i in items),has_pending=any(i.selectable for i in items),
        total=sum((i.total for i in items),Decimal('0')),pending_total=sum((i.total for i in items if i.decision=='pending'),Decimal('0')),
        paid_total=sum((i.total for i in items if i.decision=='allow'),Decimal('0')))


@login_required
def regulations(request):
    if not is_admin(request.user):raise PermissionDenied
    p=policy()
    from .forms import IntegratedPolicyForm
    form_class=IntegratedPolicyForm if settings.INTEGRATED_PURCHASES else PolicyForm
    form=form_class(request.POST if request.method=='POST' else None,instance=p)
    if request.method=='POST' and form.is_valid():
        try:form.save(request.user);messages.success(request,'구매 규정을 저장했습니다. 새 구매요청부터 적용됩니다.');return redirect('purchasing:regulations')
        except ValidationError as exc:form.add_error(None,exc)
    rows=[dict(rank=rank,mode=form[f'rank_{i}']) for i,rank in enumerate(RANKS)] if form.legacy else []
    return page(request,'regulations','구매 규정 설정',form=form,rank_rows=rows,ranks=RANKS,assigned_policy=p)


@login_required
@require_GET
def status(request):
    from .context import shell
    data=shell(request)
    response=JsonResponse(dict(user=request.user.pk,counts=data['counts']))
    add_never_cache_headers(response);return response


@login_required
@require_GET
def hints(request):
    from workflow.purchase_hints import normalize_product
    from django.db.models.functions import Replace,Lower
    from django.db.models import Value
    q=normalize_product(request.GET.get('q','')[:120]);rows=PurchaseRequest.objects.none()
    if len(q)>=2:
        rows=PurchaseRequest.objects.exclude(status__in=['draft','rejected','cancelled','denied']).filter(private=False).annotate(
            normalized=Lower(Replace(Replace('product',Value(' '),Value('')),Value('-'),Value('')))).filter(normalized__contains=q).order_by('-created_at')
    results=[dict(product=d.product,department=d.department,quantity=d.quantity,
        received=d.received_quantity,usage_department=d.usage_department,location=d.storage_location if d.received_quantity else '',requested=d.created_at.date().isoformat(),status=progress(d,simplified=True)['label']) for d in rows[:8]]
    count=rows.count()
    if settings.INTEGRATED_PURCHASES and len(q)>=2:
        from workflow.models import Document
        old=Document.objects.filter(kind__in=['office','stock'],purchase_private=False,status__in=['review','approve','executive','approved']).annotate(
            normalized=Lower(Replace(Replace('product',Value(' '),Value('')),Value('-'),Value('')))).filter(normalized__contains=q).select_related('owner__profile').order_by('-created_at')
        count+=old.count()
        results.extend(dict(product=d.product,department=d.owner.profile.department,quantity=d.quantity,received=d.quantity if d.shipment=='received' else 0,usage_department='',location='',requested=d.created_at.date().isoformat(),status='기존 문서 · '+(d.shipment_label if d.status=='approved' else d.get_status_display())) for d in old[:8])
        results=sorted(results,key=lambda d:d['requested'],reverse=True)[:8]
    response=JsonResponse(dict(count=count,results=results))
    add_never_cache_headers(response);return response


@login_required
def staff_permissions(request):
    if settings.INTEGRATED_PURCHASES:return redirect('staff')
    if not is_admin(request.user):raise PermissionDenied
    from .services import lock_policy,actor_now
    users=get_user_model().objects.filter(is_active=True,profile__approved=True,profile__representative_slot__isnull=True).select_related('profile').order_by('first_name','pk')
    if request.method=='POST':
        try:
            target=int(request.POST.get('user','0'))
            with transaction.atomic():
                p=lock_policy();actor=actor_now(request.user)
                if not is_admin(actor):raise PermissionDenied
                user=get_object_or_404(users,pk=target)
                if request.POST.get('revision')!=str(p.version):raise ValidationError('권한 설정이 바뀌었습니다. 새로고침 후 다시 저장해 주세요.')
                allowed=request.POST.get('view_completed')=='on'
                if p.department_routing and p.enabled and not allowed and target in [p.dev_director_id,p.production_director_id]:raise ValidationError('규정관리에서 담당 이사를 먼저 변경해 주세요.')
                if p.department_routing and not allowed and PurchaseRequest.objects.filter(purchase_checker=user,confirmed_at=None,status__in=['review','approval','director','ready','supplement']).exists():raise ValidationError('이 담당자에게 배정된 진행 문서가 남아 있습니다. 처리 완료 후 권한을 해제해 주세요.')
                current=PurchasePermission.objects.filter(user=user,view_completed=True).exists()
                if current and not allowed and p.enabled and not PurchasePermission.objects.exclude(user=user).filter(view_completed=True,user__is_active=True,user__profile__approved=True).exists():raise ValidationError('구매 품목확인자를 한 명 이상 남겨 주세요.')
                PurchasePermission.objects.update_or_create(user=user,defaults={'view_completed':allowed})
                p.version+=1;p.save(update_fields=['version','updated_at'])
                event(actor,'직원 구매 권한 변경',detail=f'{user.first_name}: 구매요청승인완료함 {allowed}')
            messages.success(request,'직원 구매 권한을 저장했습니다.');return redirect('purchasing:staff')
        except (ValueError,ValidationError) as exc:messages.error(request,' '.join(exc.messages) if isinstance(exc,ValidationError) else '직원을 다시 선택해 주세요.')
    permissions=set(PurchasePermission.objects.filter(view_completed=True).values_list('user_id',flat=True))
    rows=[dict(user=u,view_completed=u.pk in permissions) for u in users]
    return page(request,'staff','직원 권한관리',staff_rows=rows,revision=policy().version)


@login_required
@require_POST
def upload_attachments(request,pk):
    from .services import lock_policy,actor_now
    form=AttachmentForm(request.POST,request.FILES)
    d=get_object_or_404(requests_for(request.user),pk=pk)
    if not can_upload(request.user,d):raise PermissionDenied
    if form.is_valid():
        try:
            with StoredFiles() as stored,transaction.atomic():
                lock_policy();user=actor_now(request.user)
                d=PurchaseRequest.objects.select_for_update().get(pk=pk)
                if not requests_for(user).filter(pk=pk).exists() or not can_upload(user,d):raise PermissionDenied
                stored.add(d,user,form.cleaned_data['files'],'request' if d.owner_id==user.pk else 'purchase')
            messages.success(request,'첨부파일을 등록했습니다.')
        except ValidationError as exc:messages.error(request,' '.join(exc.messages))
    else:messages.error(request,' '.join(str(e) for errors in form.errors.values() for e in errors))
    return redirect('purchasing:detail',pk=pk)


@login_required
@require_GET
def download_attachment(request,pk):
    from .models import PurchaseAttachment
    a=get_object_or_404(PurchaseAttachment.objects.select_related('request__current_report'),pk=pk)
    if not requests_for(request.user).filter(pk=a.request_id).exists():raise Http404
    if a.kind=='purchase' and not can_view_purchase(request.user,a.request):raise Http404
    try:handle=a.file.open('rb')
    except FileNotFoundError:raise Http404('첨부파일을 찾을 수 없습니다.')
    response=FileResponse(handle,as_attachment=True,filename=a.original_name,content_type='application/octet-stream')
    response['X-Content-Type-Options']='nosniff';response['Content-Security-Policy']="default-src 'none'";add_never_cache_headers(response)
    return response


@login_required
@require_POST
def update_location(request,pk):
    from .services import lock_policy,actor_now
    with transaction.atomic():
        lock_policy();user=actor_now(request.user)
        d=get_object_or_404(PurchaseRequest.objects.select_for_update(),pk=pk,owner=user)
        if not d.received_quantity:raise PermissionDenied
        location=request.POST.get('storage_location','').strip()
        if len(location)>150:messages.error(request,'보관 위치는 150자 이내로 입력해 주세요.')
        else:
            before=d.storage_location;d.storage_location=location;d.save(update_fields=['storage_location'])
            event(user,'보관 위치 변경',request=d,detail=f'{before or "미등록"} → {location or "미등록"}')
            messages.success(request,'보관 위치를 저장했습니다.')
    return redirect('purchasing:detail',pk=pk)


@login_required
def supplement(request,pk):
    d=get_object_or_404(PurchaseRequest,pk=pk,owner=request.user,status='supplement')
    form=SupplementForm(request.POST if request.method=='POST' else None,request.FILES if request.method=='POST' else None,initial=dict(revision=d.revision,reason=d.reason,url=d.url))
    if request.method=='POST' and form.is_valid():
        try:
            with StoredFiles() as stored,transaction.atomic():
                d=supplement_request(pk,request.user,form.cleaned_data,form.cleaned_data['revision'])
                stored.add(d,request.user,form.cleaned_data['files'],'request')
            messages.success(request,'보완 내용을 제출했습니다. 구매 품목확인자가 다시 확인합니다.')
            return redirect('purchasing:detail',pk=pk)
        except ValidationError as exc:form.add_error(None,exc)
    return page(request,'supplement','구매요청 보완',doc=d,form=form,supplement=d.supplements.first())


@login_required
def amount_change(request,pk):
    d=get_object_or_404(requests_for(request.user),pk=pk)
    if not finance_assigned(request.user,d):raise PermissionDenied
    form=AmountChangeForm(request.POST if request.method=='POST' else None,initial=dict(revision=d.revision,quantity=d.quantity,unit_price=d.unit_price))
    if request.method=='POST' and form.is_valid():
        try:
            values=form.cleaned_data
            d=amend_amount(pk,request.user,values['quantity'],values['unit_price'],values['reason'],values['revision'])
            messages.success(request,'변경된 금액으로 구매요청함에서 다시 결제 요청해 주세요.' if d.current_report_id and d.current_report.payment_tracking else '금액 증가로 재승인이 필요합니다. 구매요청함에서 다시 요청해 주세요.' if d.status=='reapproval' else '금액 변경 이력을 저장했습니다. 기존 승인은 유지됩니다.')
            return redirect('purchasing:detail',pk=pk)
        except ValidationError as exc:form.add_error(None,exc)
    return page(request,'amount_change','승인 후 금액 변경',doc=d,form=form)

@login_required
@require_POST
def open_notice(request,pk):
    n=get_object_or_404(notices_for(request.user),pk=pk)
    n.read=True;n.save(update_fields=['read'])
    return redirect('purchasing:detail',pk=n.request_id) if n.request_id else redirect('purchasing:report_detail',pk=n.report_id)


def inbox_counts(user):
    qs=finance_requests(user)
    return dict(items=qs.filter(Q(status__in=['ready','reapproval'])|Q(status='hold',hold_ready=True)).filter(confirmed_at__isnull=False).count(),sent=reports_for(user).count(),hold=qs.filter(status='hold').count(),completed=qs.filter(status='received').count())


@login_required
@require_GET
def purchase_summary(request):
    from django.core.paginator import Paginator
    from django.http import QueryDict
    from .statistics import SummaryFilterForm,aggregate,workbook
    require_finance(request.user)
    today=timezone.localdate()
    scope=finance_requests(request.user).filter(status__in=['paid','received'],paid_at__isnull=False).select_related('owner')
    data=request.GET.copy()
    data.setdefault('start',today.replace(day=1).isoformat());data.setdefault('end',today.isoformat())
    form=SummaryFilterForm(data,scope=scope)
    result=dict(documents=[],total=Decimal('0'),count=0,teams=[],people=[])
    export_url=pagination_base=''
    valid=form.is_valid()
    if valid:
        result=aggregate(scope,form.cleaned_data)
        query=QueryDict(mutable=True)
        for field in ['start','end']:query[field]=form.cleaned_data[field].isoformat()
        query['department']=form.cleaned_data['department']
        query.setlist('employees',sorted(set(form.cleaned_data['employees'])))
        pagination_base=reverse('purchasing:summary')+'?'+query.urlencode()+'&page='
        export_url=reverse('purchasing:summary')+'?'+query.urlencode()+'&download=1'
        if request.GET.get('download')=='1':
            from io import BytesIO
            from django.utils.http import content_disposition_header
            names=[person['name'] for person in form.people if person['id'] in form.cleaned_data['employees']]
            output=BytesIO();workbook(result,form.cleaned_data,names).save(output)
            response=HttpResponse(output.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            filename=f'구매집계_{form.cleaned_data["start"]:%Y%m%d}_{form.cleaned_data["end"]:%Y%m%d}.xlsx'
            response['Content-Disposition']=content_disposition_header(True,filename);add_never_cache_headers(response)
            event(request.user,'구매집계 엑셀 다운로드',detail=f'{form.cleaned_data["start"]} ~ {form.cleaned_data["end"]} / {result["count"]}건 / {result["total"]}원')
            return response
    page_obj=Paginator(result['documents'],25).get_page(request.GET.get('page'))
    response=page(request,'summary','구매 집계',form=form,result=result,page_obj=page_obj,pagination_base=pagination_base,export_url=export_url,
        today=today.isoformat(),month_start=today.replace(day=1).isoformat(),year_start=today.replace(month=1,day=1).isoformat(),
        purchase_tab='summary',inbox_counts=inbox_counts(request.user),filters_valid=valid)
    if not valid and request.GET.get('download')=='1':response.status_code=400
    return response


@login_required
@require_POST
def assignments(request):
    if not settings.INTEGRATED_PURCHASES or not is_admin(request.user):raise PermissionDenied
    from .forms import PurchaseAssignmentsForm
    form=PurchaseAssignmentsForm(request.POST)
    if form.is_valid():
        try:form.save(request.user);messages.success(request,'구매 업무 담당자와 권한을 저장했습니다. 새 요청부터 적용됩니다.');return redirect('staff')
        except ValidationError as exc:form.add_error(None,exc)
    return page(request,'assignments','구매 업무 담당자·권한',form=form)


@login_required
@require_POST
def report_export(request,pk):
    from io import BytesIO
    from openpyxl import Workbook
    from django.utils.http import content_disposition_header
    report=get_object_or_404(reports_for(request.user),pk=pk)
    book=Workbook();sheet=book.active;sheet.title='구매요청'
    sheet.append(['구매요청서','원본 문서','요청 부서','요청자','품목','구매처','수량','단가','금액','구매사유'])
    for item in report.items.select_related('request'):
        s=item.snapshot
        sheet.append([report.number,s.get('number',''),s.get('department',''),s.get('owner',''),s.get('product',''),s.get('vendor',''),s.get('quantity',0),Decimal(s.get('unit_price','0')),item.total,s.get('reason','')])
        for i in [1,2,3,4,5,6,10]:sheet.cell(sheet.max_row,i).data_type='s'
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    from openpyxl.styles import Alignment,Font,PatternFill
    from openpyxl.utils import get_column_letter
    for col in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(col)].width=25 if col in [5,6,10,12,13] else 18
    for cells in sheet:
        for cell in cells:cell.alignment=Alignment(vertical='top',wrap_text=True)
    for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='0B4F3D')
    output=BytesIO();book.save(output)
    response=HttpResponse(output.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition']=content_disposition_header(True,report.number+'.xlsx');add_never_cache_headers(response)
    event(request.user,'구매요청 엑셀 다운로드',report=report)
    return response
