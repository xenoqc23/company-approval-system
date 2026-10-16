from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.comments import Comment
from .leave_records import leave_usage_summaries


def purchase_workbook(documents):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '구매요청서'
    sheet.append(['성명', '품명', '수량', '단가', '총가격', '신청사유'])
    for doc in documents:
        sheet.append([doc.owner.first_name, doc.product, doc.quantity,
                      doc.unit_price, doc.total_price, doc.reason])
        # User-entered text remains literal, even when it starts with '='.
        for column in [1, 2, 6]:
            sheet.cell(sheet.max_row, column).data_type = 's'
        for column in [4, 5]:
            sheet.cell(sheet.max_row, column).number_format = '#,##0'
        for cell in sheet[sheet.max_row]:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='315D43')
        cell.font = Font(color='FFFFFF', bold=True)
    for column, width in [('A', 16), ('B', 32), ('C', 10), ('D', 18), ('E', 20), ('F', 55)]:
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    return workbook


def leave_balances_workbook(rows, year):
    rows = list(rows)
    histories = leave_usage_summaries([row['user'] for row in rows], year)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f'{year}년 연차 현황'
    sheet.append(['연도', '성명', '부서', '직급', '총연차', '사용연차', '잔여연차', '공용연차', '사용내역'])
    for row in rows:
        person = row['user']
        sheet.append([year, person.first_name or person.username, person.profile.department,
                      person.profile.rank, row['total'], row['used'], row['remaining'],
                      row['common'], histories.get(person.pk) or '—'])
        for column in [2, 3, 4, 9]:
            sheet.cell(sheet.max_row, column).data_type = 's'
        for column in [5, 6, 7, 8]:
            cell = sheet.cell(sheet.max_row, column)
            cell.number_format = '#,##0' if cell.value == int(cell.value) else '#,##0.##'
            if cell.value < 0:
                cell.font = Font(color='A13F31')
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='315D43')
        cell.font = Font(color='FFFFFF', bold=True)
    for column in 'ABCDEFGH':
        sheet.column_dimensions[column].width = 18
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    overview = workbook.create_sheet('인원별 연차표')
    overview.append(['성명', '근무연수', '연차일수', '공용연차', '사용가능연차', '잔여연차', '사용내역'])
    for row in rows:
        overview.append([row['user'].first_name or row['user'].username, row['tenure'],
                         row['total'], row['common'], row['available'], row['remaining'],
                         histories.get(row['user'].pk) or '—'])
        overview.cell(overview.max_row, 1).data_type = 's'
        overview.cell(overview.max_row, 7).data_type = 's'
        for cell in list(overview[overview.max_row])[2:6]:
            cell.number_format = '#,##0' if cell.value == int(cell.value) else '#,##0.##'
            if cell.value < 0:
                cell.font = Font(color='A13F31')
    for cell in overview[1]:
        cell.fill = PatternFill('solid', fgColor='343F57')
        cell.font = Font(color='FFFFFF', bold=True)
    for column in 'ABCDEF':
        overview.column_dimensions[column].width = 20
    overview.freeze_panes = 'A2'
    overview.auto_filter.ref = overview.dimensions
    for target, history_column in [(sheet, 'I'), (overview, 'G')]:
        target.column_dimensions[history_column].width = 72
        target[f'{history_column}1'].comment = Comment(
            '조회 연도의 결재 완료 휴가·경리 직접 등록·공용연차를 날짜 내림차순으로 표시합니다. '
            '승인·등록된 사용 예정분도 포함하며 취소·반려·대기 내역은 제외합니다. '
            '날짜를 알 수 없는 기존 잔액 정정은 사용내역에 표시하지 않습니다.', 'Xenosys')
        for number in range(2, target.max_row + 1):
            cell = target[f'{history_column}{number}']
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            # Excel's column width is approximate; allow room for Korean labels.
            text_width = sum(2 if ord(char) > 127 else 1 for char in str(cell.value))
            target.row_dimensions[number].height = min(409, max(30, 16 * (1 + text_width // 65)))
    return workbook
