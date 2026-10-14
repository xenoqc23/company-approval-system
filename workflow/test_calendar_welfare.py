from datetime import date, time
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.test import TestCase, Client
from django.urls import reverse
from openpyxl import load_workbook
from .annual_grants import entitlement, sync_annual_grants, set_hire_date
from .calendars import month_grid
from .company_holidays import save_company_holiday, delete_company_holiday
from .exports import leave_balances_workbook
from .leave_records import balance_rows
from .models import AnnualBalance, Absence, Audit, CompanyHoliday, CompanyLeaveCharge, Document, Profile
from .services import save_document, act_on_document, edit_balance
from .welfare import award_candidates, BENEFITS


class CalendarWelfareTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        User = get_user_model()
        for key, flags in [('employee', {}), ('other', {}), ('accounting', {'manage_leave': True}),
                           ('reviewer', {}), ('approver', {}), ('admin', {'manage_system': True}),
                           ('pending', {'approved': False}), ('inactive', {})]:
            user = User.objects.create_user(key, first_name=key, password='1234', is_active=key != 'inactive')
            approved = flags.pop('approved', True)
            Profile.objects.create(user=user, approved=approved, department='개발팀', rank='대리', **flags)
            AnnualBalance.objects.create(user=user, year=2026, total=15, used=Decimal('0.25'), grant_initialized=True)
            setattr(cls, key, user)

    def holiday(self, **overrides):
        values = {'title': '회사 지정 휴무', 'day': date(2026, 10, 23), 'use_annual': True, 'description': ''}
        values.update(overrides)
        return save_company_holiday(self.accounting, values)

    def finish_leave(self, **overrides):
        values = {'owner': self.employee, 'reviewer': self.reviewer, 'approver': self.approver,
                  'kind': 'leave', 'start_date': date(2026, 10, 19), 'end_date': date(2026, 10, 19),
                  'reason': 'private reason should not leak'}
        values.update(overrides)
        doc = save_document(Document(**values), self.employee, submit=True)
        act_on_document(doc.pk, self.reviewer, 'review')
        return act_on_document(doc.pk, self.approver, 'approve')

    def test_entitlement_uses_january_first_and_caps_25(self):
        hire = date(2023, 10, 15)
        self.assertEqual(entitlement(hire, 2026), 15)
        self.assertEqual(entitlement(hire, 2027), 16)
        self.assertEqual(entitlement(date(2026, 10, 1), 2026), 11)
        self.assertEqual(entitlement(date(1990, 1, 1), 2026), 25)
        for years, days in [(0, 11), (1, 15), (2, 15), (3, 16), (4, 16), (5, 17), (21, 25), (40, 25)]:
            self.assertEqual(entitlement(date(2026-years, 1, 1), 2026), days)

    def test_rollover_preserves_preapproved_usage_no_carry_no_repeat(self):
        self.employee.profile.hire_date = date(2023, 10, 15)
        self.employee.profile.save()
        AnnualBalance.objects.create(user=self.employee, year=2027, used=Decimal('1.25'))
        self.assertEqual(sync_annual_grants(date(2027, 1, 1)), 1)
        balance = AnnualBalance.objects.get(user=self.employee, year=2027)
        self.assertEqual((balance.total, balance.used, balance.remaining), (16, Decimal('1.25'), Decimal('14.75')))
        edit_balance(self.accounting, self.employee, 2027, 'total', Decimal('20'), '개별 조정')
        self.assertEqual(sync_annual_grants(date(2027, 10, 15)), 0)
        self.assertEqual(AnnualBalance.objects.get(user=self.employee, year=2027).total, 20)
        self.assertEqual(AnnualBalance.objects.get(user=self.employee, year=2026).total, 15)

    def test_hire_save_does_not_overwrite_existing_manual_total_apply_preserves_used(self):
        set_hire_date(self.accounting, self.employee, date(2023, 10, 15), 2026)
        sync_annual_grants(date(2026, 10, 16))
        self.assertEqual(AnnualBalance.objects.get(user=self.employee, year=2026).total, 15)
        set_hire_date(self.accounting, self.employee, date(2000, 1, 1), 2026, apply=True)
        balance = AnnualBalance.objects.get(user=self.employee, year=2026)
        self.assertEqual((balance.total, balance.used), (25, Decimal('0.25')))
        with self.assertRaises(PermissionDenied):
            set_hire_date(self.other, self.employee, date(2000, 1, 1), 2026, apply=True)
        with self.assertRaises(ValidationError):
            set_hire_date(self.accounting, self.employee, date(2099, 1, 1), 2026)

    def test_company_leave_applies_only_active_approved_humans_immediately(self):
        holiday = self.holiday()
        self.assertEqual(set(holiday.charges.values_list('employee_id', flat=True)),
                         {self.employee.pk, self.other.pk, self.accounting.pk, self.reviewer.pk, self.approver.pk})
        for person in [self.employee, self.accounting, self.other]:
            self.assertEqual(AnnualBalance.objects.get(user=person, year=2026).used, Decimal('1.25'))
        for person in [self.admin, self.pending, self.inactive]:
            self.assertEqual(AnnualBalance.objects.get(user=person, year=2026).used, Decimal('0.25'))

    def test_company_ordinary_holiday_has_no_charges(self):
        holiday = self.holiday(use_annual=False)
        self.assertEqual(holiday.charges.count(), 0)
        self.assertEqual(AnnualBalance.objects.get(user=self.employee, year=2026).used, Decimal('0.25'))

    def test_company_delete_restores_once_original_people_year_even_if_inactive(self):
        holiday = self.holiday(day=date(2027, 1, 8))
        self.employee.is_active = False
        self.employee.save()
        delete_company_holiday(self.accounting, holiday.pk)
        self.assertEqual(AnnualBalance.objects.get(user=self.employee, year=2027).used, 0)
        self.assertEqual(AnnualBalance.objects.get(user=self.employee, year=2026).used, Decimal('0.25'))
        self.assertFalse(CompanyLeaveCharge.objects.filter(holiday=holiday, active=True).exists())
        with self.assertRaises(CompanyHoliday.DoesNotExist):
            delete_company_holiday(self.accounting, holiday.pk)
        self.assertTrue(Audit.objects.filter(target=self.employee, event='2027년 공용연차 복원').exists())

    def test_company_date_change_and_type_change_restore_then_apply(self):
        holiday = self.holiday()
        values = {'title': '새 제목', 'day': date(2027, 1, 8), 'use_annual': True, 'description': ''}
        save_company_holiday(self.accounting, values, holiday.pk)
        self.assertEqual(AnnualBalance.objects.get(user=self.employee, year=2026).used, Decimal('0.25'))
        self.assertEqual(AnnualBalance.objects.get(user=self.employee, year=2027).used, 1)
        values['use_annual'] = False
        save_company_holiday(self.accounting, values, holiday.pk)
        self.assertEqual(AnnualBalance.objects.get(user=self.employee, year=2027).used, 0)

    def test_company_title_edit_does_not_charge_again_or_add_new_employee(self):
        holiday = self.holiday()
        self.pending.profile.approved = True
        self.pending.profile.save()
        values = {'title': '제목 수정', 'day': holiday.day, 'use_annual': True, 'description': '설명'}
        save_company_holiday(self.accounting, values, holiday.pk)
        self.assertEqual(AnnualBalance.objects.get(user=self.employee, year=2026).used, Decimal('1.25'))
        self.assertFalse(holiday.charges.filter(employee=self.pending).exists())
        with self.assertRaises(ValidationError):
            self.holiday()
        with self.assertRaises(PermissionDenied):
            save_company_holiday(self.other, values)

    def test_calendar_shows_approved_only_cancel_removes_and_detail_hides_private_reason(self):
        doc = self.finish_leave()
        self.client.force_login(self.other)
        response = self.client.get(reverse('calendar'), {'year': 2026, 'month': 10})
        self.assertContains(response, 'employee · 연차')
        detail = self.client.get(reverse('calendar_leave', args=[doc.pk]))
        self.assertEqual(detail.status_code, 200)
        self.assertNotContains(detail, 'private reason')
        self.assertNotContains(detail, '휴가원 상세 보기')
        self.assertEqual(self.client.get(reverse('detail', args=[doc.pk])).status_code, 404)
        act_on_document(doc.pk, self.approver, 'cancel_approval', '변경')
        self.assertNotContains(self.client.get(reverse('calendar'), {'year': 2026, 'month': 10}), 'employee · 연차')
        self.assertEqual(self.client.get(reverse('calendar_leave', args=[doc.pk])).status_code, 404)

    def test_public_substitute_holidays_and_weekends_display(self):
        grid = month_grid(2026, 10, self.other)
        days = {cell['date']: cell for week in grid['weeks'] for cell in week}
        self.assertIn('한글날', days[date(2026, 10, 9)]['holidays'])
        self.assertTrue(any('대체공휴일' in label for label in days[date(2026, 10, 5)]['holidays']))
        doc = self.finish_leave(start_date=date(2026, 10, 9), end_date=date(2026, 10, 12))
        grid = month_grid(2026, 10, self.other)
        events = {cell['date']: cell['events'] for week in grid['weeks'] for cell in week}
        self.assertTrue(events[date(2026, 10, 9)])
        self.assertFalse(events[date(2026, 10, 10)])
        self.assertEqual(doc.charged, 2)  # Holiday calendar does not change agreed leave calculation.

    def test_any_employee_edits_deletes_absence_no_leave_charge_no_audit(self):
        self.client.force_login(self.employee)
        payload = {'reason': '예비군', 'start_date': '2026-10-20', 'end_date': '2026-10-20', 'all_day': 'on', 'employee': self.other.pk}
        before = Audit.objects.count()
        self.assertEqual(self.client.post(reverse('absence_new'), payload).status_code, 302)
        absence = Absence.objects.get()
        self.assertEqual(absence.employee, self.employee)
        self.client.force_login(self.other)
        payload['reason'] = '교육'
        self.assertEqual(self.client.post(reverse('absence_edit', args=[absence.pk]), payload).status_code, 302)
        absence.refresh_from_db()
        self.assertEqual((absence.employee, absence.reason), (self.employee, '교육'))
        self.assertEqual(self.client.get(reverse('absence_delete', args=[absence.pk])).status_code, 405)
        self.assertEqual(self.client.post(reverse('absence_delete', args=[absence.pk])).status_code, 302)
        self.assertEqual(Absence.objects.count(), 0)
        self.assertEqual(Audit.objects.count(), before)
        self.assertEqual(AnnualBalance.objects.get(user=self.employee, year=2026).used, Decimal('0.25'))

    def test_invalid_absence_date_time_and_csrf(self):
        self.client.force_login(self.other)
        payload = {'reason': '외근', 'start_date': '2026-10-20', 'end_date': '2026-10-19', 'all_day': 'on'}
        self.assertEqual(self.client.post(reverse('absence_new'), payload).status_code, 200)
        self.assertFalse(Absence.objects.exists())
        payload.update(end_date='2026-10-20', all_day='', start_time='12:00', end_time='11:00')
        self.client.post(reverse('absence_new'), payload)
        self.assertFalse(Absence.objects.exists())
        secure = Client(enforce_csrf_checks=True)
        secure.force_login(self.other)
        self.assertEqual(secure.post(reverse('absence_new'), payload).status_code, 403)

    def test_company_manager_routes_reject_regular_users_and_pending_calendar(self):
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(reverse('company_holidays')).status_code, 403)
        self.assertEqual(self.client.post(reverse('company_holidays'), {}).status_code, 403)
        holiday = self.holiday()
        self.assertEqual(self.client.post(reverse('company_holiday_delete', args=[holiday.pk])).status_code, 403)
        self.client.force_login(self.pending)
        self.assertEqual(self.client.get(reverse('calendar')).url, reverse('pending'))

    def test_awards_actual_anniversary_all_milestones_leap_day_and_admin_exclusion(self):
        self.employee.profile.hire_date = date(2023, 10, 15)
        self.employee.profile.save()
        self.admin.profile.hire_date = date(2023, 10, 15)
        self.admin.profile.save()
        candidates = award_candidates(2026, date(2026, 10, 14))
        self.assertEqual([(row['employee'].pk, row['day'], row['reward'], row['arrived']) for row in candidates],
                         [(self.employee.pk, date(2026, 10, 15), '금 3돈', False)])
        self.assertTrue(award_candidates(2026, date(2026, 10, 15))[0]['arrived'])
        self.assertFalse(award_candidates(2027, date(2027, 1, 1)))
        self.employee.profile.hire_date = date(2024, 2, 29)
        self.employee.profile.save()
        self.assertEqual(award_candidates(2027, date(2027, 1, 1))[0]['day'], date(2027, 2, 28))

    def test_welfare_tabs_render_benefit_values_and_leave_table(self):
        self.client.force_login(self.other)
        for tab in ['awards', 'benefits', 'leave']:
            self.assertEqual(self.client.get(reverse('welfare'), {'tab': tab}).status_code, 200)
        response = self.client.get(reverse('welfare'), {'tab': 'benefits'})
        self.assertContains(response, '부모사망(배우자 포함)')
        self.assertContains(response, '조부모사망(배우자 포함)')
        self.assertContains(response, '2,500,000원')
        self.assertEqual(len(BENEFITS), 22)

    def test_company_table_excel_negative_and_fractional_usage(self):
        AnnualBalance.objects.filter(user=self.employee, year=2026).update(total=0)
        self.holiday()
        rows = balance_rows(2026)
        row = next(r for r in rows if r['user'] == self.employee)
        self.assertEqual((row['common'], row['available'], row['remaining']), (1, -1, Decimal('-1.25')))
        output = BytesIO()
        leave_balances_workbook(rows, 2026).save(output)
        book = load_workbook(BytesIO(output.getvalue()))
        self.assertEqual(list(book['인원별 연차표'].values)[0], ('성명', '근무연수', '연차일수', '공통사용연차', '사용가능연차', '잔여연차'))

    def test_all_new_templates_and_calendar_bad_parameters_render(self):
        self.client.force_login(self.accounting)
        for name in ['calendar', 'absence_new', 'company_holidays', 'balances', 'welfare']:
            self.assertEqual(self.client.get(reverse(name)).status_code, 200)
        self.assertEqual(self.client.get(reverse('calendar'), {'year': 'bad', 'month': '13'}).status_code, 200)

    def test_preapproved_next_year_leave_and_rollover_then_cancellation(self):
        self.employee.profile.hire_date = date(2023, 10, 15)
        self.employee.profile.save()
        doc = self.finish_leave(start_date=date(2027, 1, 4), end_date=date(2027, 1, 4))
        before = AnnualBalance.objects.get(user=self.employee, year=2027)
        self.assertEqual((before.total, before.used), (16, 1))
        sync_annual_grants(date(2027, 1, 1))
        act_on_document(doc.pk, self.approver, 'cancel_approval', '취소')
        self.assertEqual(AnnualBalance.objects.get(user=self.employee, year=2027).remaining, 16)

    def test_company_apply_failure_rolls_back_every_employee(self):
        with patch('workflow.services.audit', side_effect=RuntimeError('simulated failure')):
            with self.assertRaises(RuntimeError):
                self.holiday()
        self.assertFalse(CompanyHoliday.objects.exists())
        self.assertFalse(CompanyLeaveCharge.objects.exists())
        self.assertEqual(AnnualBalance.objects.get(user=self.employee, year=2026).used, Decimal('0.25'))

    def test_company_delete_uses_post_csrf_and_ordinary_to_common_edit(self):
        holiday = self.holiday(use_annual=False)
        self.client.force_login(self.accounting)
        self.assertEqual(self.client.get(reverse('company_holiday_delete', args=[holiday.pk])).status_code, 405)
        secure = Client(enforce_csrf_checks=True)
        secure.force_login(self.accounting)
        self.assertEqual(secure.post(reverse('company_holiday_delete', args=[holiday.pk])).status_code, 403)
        save_company_holiday(self.accounting, {'title': holiday.title, 'day': holiday.day,
            'use_annual': True, 'description': ''}, holiday.pk)
        self.assertEqual(AnnualBalance.objects.get(user=self.employee, year=2026).used, Decimal('1.25'))

    def test_new_hire_form_is_manager_only_and_preserves_used(self):
        payload = {'operation': 'hire', 'user': self.employee.pk, 'year': '2026',
                   'hire_date': '2000-01-01', 'apply': 'yes'}
        self.client.force_login(self.other)
        self.assertEqual(self.client.post(reverse('balances'), payload).status_code, 403)
        self.client.force_login(self.accounting)
        self.assertEqual(self.client.post(reverse('balances'), payload).status_code, 302)
        balance = AnnualBalance.objects.get(user=self.employee, year=2026)
        self.assertEqual((balance.total, balance.used), (25, Decimal('0.25')))
