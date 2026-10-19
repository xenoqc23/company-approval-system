from datetime import date
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from .exports import leave_balances_workbook
from .leave_records import balance_rows, leave_usage_summaries
from .manual_leave import register_manual_leave, cancel_manual_leave
from .models import AnnualBalance, CompanyHoliday, CompanyLeaveCharge, Document, Profile


class LeaveExportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        from .models import ApprovalFlowPolicy
        ApprovalFlowPolicy.objects.update_or_create(pk=1, defaults={'enabled':False})
        for name, department in [('employee','품질팀'), ('finance','인사팀')]:
            user = get_user_model().objects.create_user(name, first_name=name)
            Profile.objects.create(user=user, department=department, rank='대리', approved=True,
                                   manage_leave=name == 'finance', hire_date=date(2020,1,1))
            setattr(cls, name, user)
        AnnualBalance.objects.create(user=cls.employee, year=2026, total=15, used=Decimal('3.25'),
                                     grant_initialized=True)
        AnnualBalance.objects.create(user=cls.finance, year=2026, total=17, used=0, grant_initialized=True)

    def document(self, day, leave_type='annual', **kwargs):
        fields = dict(owner=self.employee, kind='leave', leave_type=leave_type, start_date=day, end_date=day,
                      status='approved', allocations={str(day.year):'1' if leave_type == 'annual' else '.25' if leave_type == 'outing' else '.5'},
                      reason='PRIVATE_LEAVE_REASON')
        fields.update(kwargs)
        return Document.objects.create(**fields)

    def manual(self, day):
        return register_manual_leave(self.finance, self.employee, {
            'used_date':day, 'leave_type':'annual', 'reason':'PRIVATE_MANUAL_REASON'}, uuid4())[0]

    def common(self, day, active=True, amount=Decimal('1')):
        holiday = CompanyHoliday.objects.create(title='회사 휴무', day=day, use_annual=True, created_by=self.finance)
        return CompanyLeaveCharge.objects.create(holiday=holiday, employee=self.employee,
            year=day.year, amount=amount, active=active)

    def test_short_labels_sorted_dates_including_manual_common_and_excluding_cancelled(self):
        self.document(date(2026,10,16))
        self.document(date(2026,9,1), 'am')
        self.document(date(2026,9,2), 'pm')
        self.document(date(2026,8,28), 'outing')
        self.manual(date(2026,10,8))
        self.common(date(2026,10,9))
        for state in ['cancelled','rejected','review','approve','draft','deleted']:
            self.document(date(2026,10,12), status=state)
        self.document(date(2026,10,13), kind='office')
        self.document(date(2026,10,14), owner=self.finance)
        self.document(date(2026,10,15), allocations={})
        self.common(date(2026,10,5), active=False)
        cancelled = self.manual(date(2026,10,7))
        cancel_manual_leave(self.finance, self.employee, cancelled.pk, 'PRIVATE_CANCEL_REASON')
        summary = leave_usage_summaries([self.employee], 2026)[self.employee.pk]
        self.assertEqual(summary, {'personal':'10/16(연차), 10/8(연차), 9/2(반차), 9/1(반차), 8/28(외출)', 'common':'10/9(공용연차)'})

    def test_multi_day_and_year_boundary_match_weekday_charges(self):
        self.document(date(2026,12,31), end_date=date(2027,1,4), allocations={'2026':'1','2027':'2'})
        self.manual(date(2027,1,5))
        self.common(date(2027,1,6))
        self.assertEqual(leave_usage_summaries([self.employee],2026)[self.employee.pk], {'personal':'12/31(연차)', 'common':''})
        self.assertEqual(leave_usage_summaries([self.employee],2027)[self.employee.pk],
            {'personal':'1/5(연차), 1/4(연차), 1/1(연차)', 'common':'1/6(공용연차)'})
        self.document(date(2026,10,16), end_date=date(2026,10,19), allocations={'2026':'2'})
        self.assertEqual(leave_usage_summaries([self.employee],2026)[self.employee.pk],
            {'personal':'12/31(연차), 10/19(연차), 10/16(연차)', 'common':''})

    def test_real_download_both_sheets_numbers_wrap_filters_and_read_only_balances(self):
        self.document(date(2026,10,16))
        self.manual(date(2026,10,8))
        self.common(date(2026,10,9), amount=Decimal('.5'))
        before = list(AnnualBalance.objects.values())
        self.client.force_login(self.finance)
        response = self.client.post(reverse('export_balances'), {'year':2026, 'department':'품질팀'})
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        main, overview = workbook.worksheets
        summary = '10/16(연차), 10/8(연차)'
        self.assertEqual(main.max_row, 2)
        self.assertEqual(main['H1'].value, '공용연차')
        self.assertEqual(main['H2'].value, .5)
        self.assertEqual(main['H2'].number_format, '#,##0.##')
        self.assertEqual(main['I1'].value, '개인 사용내역')
        self.assertEqual(main['I2'].value, summary)
        self.assertEqual(main['J1'].value, '공용연차 내역')
        self.assertEqual(main['J2'].value, '10/9(공용연차)')
        self.assertEqual(overview['D1'].value, '공용연차')
        self.assertEqual(overview['D2'].value, .5)
        self.assertEqual(overview['G1'].value, '개인 사용내역')
        self.assertEqual(overview['G2'].value, summary)
        self.assertEqual(overview['H1'].value, '공용연차 내역')
        self.assertEqual(overview['H2'].value, '10/9(공용연차)')
        for sheet, column, last in [(main,'I','J'), (overview,'G','H')]:
            self.assertEqual(sheet[f'{column}2'].data_type, 's')
            self.assertTrue(sheet[f'{column}2'].alignment.wrap_text)
            self.assertEqual(sheet.freeze_panes, 'A2')
            self.assertEqual(sheet.max_column, 10 if sheet == main else 8)
            self.assertTrue(sheet[f'{last}2'].alignment.wrap_text)
            self.assertEqual(sheet.auto_filter.ref, f'A1:{last}2')
            self.assertIn('사용 예정', sheet[f'{column}1'].comment.text)
        all_text = ' '.join(str(cell.value) for sheet in workbook for row in sheet for cell in row)
        self.assertNotIn('PRIVATE_', all_text)
        self.assertEqual(list(AnnualBalance.objects.values()), before)

    def test_no_fabricated_usage_for_old_balance_adjustments_and_no_rows(self):
        histories = leave_usage_summaries([self.employee],2026)
        self.assertEqual(histories[self.employee.pk], {'personal':'', 'common':''})
        workbook = leave_balances_workbook(balance_rows(2026,'품질팀'),2026)
        self.assertEqual(workbook.active['I2'].value, '—')
        self.assertEqual(workbook.active['F2'].value, Decimal('3.25'))
        empty = leave_balances_workbook([],2026)
        self.assertEqual(empty.active.max_row, 1)
        self.assertEqual(empty['인원별 연차표'].max_row, 1)

    def test_export_history_uses_three_batched_queries_independent_of_employee_count(self):
        self.document(date(2026,10,16))
        with self.assertNumQueries(3):
            histories = leave_usage_summaries([self.employee,self.finance],2026)
        self.assertEqual(histories[self.employee.pk], {'personal':'10/16(연차)', 'common':''})
        self.assertEqual(histories[self.finance.pk], {'personal':'', 'common':''})
        with self.assertNumQueries(0):
            self.assertEqual(leave_usage_summaries([],2026), {})

    def test_common_only_leave_stays_out_of_personal_history_on_both_sheets(self):
        self.common(date(2026,5,4))
        workbook = leave_balances_workbook(balance_rows(2026,'품질팀'),2026)
        for sheet, personal, common in [(workbook.active,'I','J'), (workbook['인원별 연차표'],'G','H')]:
            self.assertEqual(sheet[f'{personal}2'].value, '—')
            self.assertEqual(sheet[f'{common}2'].value, '5/4(공용연차)')
