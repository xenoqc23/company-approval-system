from datetime import date, time
from decimal import Decimal
from io import BytesIO
from openpyxl import load_workbook
from urllib.parse import urlencode
from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.test import TestCase, Client
from django.urls import reverse
from workflow.models import AccountingReadState, AnnualBalance, Audit, Document, Notice, Policy, Profile
from workflow.services import act_on_document, edit_balance, save_document, visible_documents
from workflow.forms import DocumentForm
from workflow.management.commands.seed_demo import TENURE, POLICY


class ApprovalTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        User = get_user_model()
        for key, flags in [
            ('owner', {}), ('reviewer', {}), ('approver', {}), ('stranger', {}),
            ('recipient', {}), ('accounting', {'manage_leave': True, 'view_accounting': True, 'procure': True}),
            ('admin', {'manage_system': True}),
        ]:
            user = User.objects.create_user(key, first_name=key, password='test-password-only')
            Profile.objects.create(user=user, department='개발팀', rank='대리', approved=True, **flags)
            setattr(cls, key, user)
        AnnualBalance.objects.create(user=cls.owner, year=2026, total=15, used=3)
        Policy.objects.create(tenure=TENURE, text=POLICY)

    def leave(self, submit=True, **fields):
        defaults = dict(owner=self.owner, reviewer=self.reviewer, approver=self.approver,
                        kind='leave', leave_type='annual', start_date=date(2026, 10, 12),
                        end_date=date(2026, 10, 14), reason='개인 일정')
        defaults.update(fields)
        return save_document(Document(**defaults), defaults['owner'], submit=submit)

    def purchase(self, **fields):
        defaults = dict(owner=self.owner, reviewer=self.reviewer, approver=self.approver,
                        recipient=self.recipient, kind='office', product='A4 용지', quantity=4,
                        unit_price=Decimal('6500.25'), url='https://example.com/a4')
        defaults.update(fields)
        return save_document(Document(**defaults), defaults['owner'], submit=True)

    def finish(self, doc):
        reviewed = act_on_document(doc.pk, doc.reviewer, 'review')
        if doc.kind == 'stock':
            return reviewed
        return act_on_document(doc.pk, doc.approver, 'approve')

    def used(self, year=2026):
        return AnnualBalance.objects.get(user=self.owner, year=year).used

    def test_rank_order_across_people_and_document_lists(self):
        for user, rank in [(self.owner, '사원'), (self.reviewer, '팀장'), (self.approver, '이사')]:
            user.profile.rank = rank
            user.profile.save()
        from workflow.people import approval_people, rank_order
        self.assertEqual(list(rank_order(get_user_model().objects.filter(
            pk__in=[self.owner.pk, self.reviewer.pk, self.approver.pk])).values_list('pk', flat=True)),
            [self.approver.pk, self.reviewer.pk, self.owner.pk])
        self.assertEqual(list(approval_people().values_list('pk', flat=True))[:2], [self.approver.pk, self.reviewer.pk])
        form = DocumentForm()
        self.assertEqual(list(form.fields['recipient'].queryset.values_list('pk', flat=True))[:2], [self.approver.pk, self.reviewer.pk])
        low = self.finish(self.purchase())
        high = self.finish(self.purchase(owner=self.approver))
        self.client.force_login(self.accounting)
        response = self.client.get(reverse('documents'), {'mode': 'accounting', 'kind': 'office'})
        self.assertEqual([doc.pk for doc in response.context['documents']], [high.pk, low.pk])

    def test_balance_department_filter_preserves_selection_and_private_access(self):
        self.owner.profile.department = '품질팀'
        self.owner.profile.save()
        self.client.force_login(self.accounting)
        response = self.client.get(reverse('balances'), {'year': 2026, 'department': '품질팀'})
        self.assertEqual([row['user'].pk for row in response.context['rows']], [self.owner.pk])
        response = self.client.post(reverse('balances') + '?year=2026&department=품질팀',
            {'user': self.owner.pk, 'year': 2026, 'field': 'total', 'value': '16', 'reason': '근속 변경'})
        self.assertEqual(response.status_code, 302)
        response = self.client.get(response.url)
        self.assertEqual(response.context['selected_department'], '품질팀')
        self.assertEqual(response.context['rows'][0]['total'], 16)
        self.client.force_login(self.stranger)
        response = self.client.get(reverse('balances'), {'department': '품질팀'})
        self.assertEqual([row['user'].pk for row in response.context['rows']], [self.stranger.pk])
        self.assertNotContains(response, 'name="department"')

    def test_purchase_excel_selected_rows_prices_and_literal_text(self):
        selected = self.finish(self.purchase(product='=SUM(1,2)', reason='=HYPERLINK("https://example.com")'))
        self.finish(self.purchase(product='미선택 품목'))
        self.client.force_login(self.accounting)
        response = self.client.post(reverse('export_purchases'), {'documents': [selected.pk, selected.pk]})
        self.assertEqual(response.status_code, 200)
        self.assertIn('attachment;', response['Content-Disposition'])
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(list(sheet.values), [
            ('성명', '품명', '수량', '단가', '총가격', '신청사유'),
            ('owner', selected.product, 4, 6500.25, 26001, selected.reason),
        ])
        self.assertEqual(sheet['B2'].data_type, 's')
        self.assertEqual(sheet['F2'].data_type, 's')
        self.assertEqual(sheet['D2'].data_type, 'n')
        self.assertEqual(sheet['D2'].number_format, '#,##0')
        self.assertEqual(sheet['E2'].number_format, '#,##0')
        self.assertTrue(Audit.objects.filter(event='구매요청서 엑셀 다운로드', actor=self.accounting).exists())
        selected.refresh_from_db()
        self.assertEqual(selected.shipment, '')

    def test_purchase_excel_refuses_unapproved_stock_cancelled_missing_and_empty(self):
        valid = self.finish(self.purchase())
        pending = self.purchase()
        stock = self.finish(self.purchase(kind='stock'))
        cancelled = self.finish(self.purchase())
        act_on_document(cancelled.pk, self.approver, 'cancel_approval', '취소')
        self.client.force_login(self.accounting)
        for ids in [[], ['x'], ['١'], ['9' * 30], [valid.pk] * 101,
                    [valid.pk, pending.pk], [stock.pk], [cancelled.pk], [valid.pk, 999999]]:
            with self.subTest(ids=ids[:3]):
                self.assertEqual(self.client.post(reverse('export_purchases'), {'documents': ids}).status_code, 400)
        self.assertFalse(Audit.objects.filter(event='구매요청서 엑셀 다운로드').exists())

    def test_purchase_excel_requires_accounting_permission_post_and_csrf(self):
        doc = self.finish(self.purchase())
        for user in [self.owner, self.admin]:
            self.client.force_login(user)
            self.assertEqual(self.client.post(reverse('export_purchases'), {'documents': [doc.pk]}).status_code, 403)
        self.client.force_login(self.accounting)
        self.assertEqual(self.client.get(reverse('export_purchases')).status_code, 405)
        secure_client = Client(enforce_csrf_checks=True)
        secure_client.force_login(self.accounting)
        self.assertEqual(secure_client.post(reverse('export_purchases'), {'documents': [doc.pk]}).status_code, 403)

    def test_only_final_approval_charges_and_only_final_actor_cancels_once(self):
        doc = self.leave()
        self.assertEqual(self.used(), 3)
        act_on_document(doc.pk, self.reviewer, 'review')
        self.assertEqual(self.used(), 3)
        act_on_document(doc.pk, self.approver, 'approve')
        self.assertEqual(self.used(), 6)
        for user, action in [(self.reviewer, 'cancel_review'), (self.admin, 'cancel_approval'),
                             (self.approver, 'approve')]:
            with self.assertRaises(PermissionDenied):
                act_on_document(doc.pk, user, action, '사유')
        act_on_document(doc.pk, self.approver, 'cancel_approval', '일정 취소')
        self.assertEqual(self.used(), 3)
        with self.assertRaises(PermissionDenied):
            act_on_document(doc.pk, self.approver, 'cancel_approval', '중복 요청')
        self.assertEqual(self.used(), 3)
        self.assertTrue(Audit.objects.filter(document=doc, event='최종 승인 취소').exists())

    def test_year_boundary_cancellation_restores_original_buckets(self):
        AnnualBalance.objects.create(user=self.owner, year=2027, total=15, used=2)
        doc = self.finish(self.leave(start_date=date(2026, 12, 31), end_date=date(2027, 1, 4)))
        # No public-holiday calendar: Fri Jan 1 is counted; weekend Jan 2/3 excluded.
        self.assertEqual(doc.allocations, {'2026': '1', '2027': '2'})
        self.assertEqual(self.used(2026), 4)
        self.assertEqual(self.used(2027), 4)
        edit_balance(self.accounting, self.owner, 2026, 'total', Decimal('20'), '수동 갱신')
        act_on_document(doc.pk, self.approver, 'cancel_approval', '취소')
        self.assertEqual(self.used(2026), 3)
        self.assertEqual(self.used(2027), 2)

    def test_negative_leave_is_allowed(self):
        b = AnnualBalance.objects.get(user=self.owner, year=2026)
        b.total = 1
        b.save()
        self.finish(self.leave())
        b.refresh_from_db()
        self.assertEqual(b.remaining, Decimal('-5'))

    def test_fixed_half_and_outing_deduction_independent_of_duration(self):
        for i, (kind, amount) in enumerate([('am', '.5'), ('pm', '.5'), ('outing', '.25')]):
            day = date(2026, 10, 12 + i)
            doc = self.finish(self.leave(leave_type=kind, start_date=day, end_date=day,
                                         start_time=time(9), end_time=time(12, 30)))
            self.assertEqual(doc.charged, Decimal(amount))
        self.assertEqual(self.used(), Decimal('4.25'))

    def test_weekends_excluded_and_weekend_only_rejected(self):
        doc = self.finish(self.leave(start_date=date(2026, 10, 9), end_date=date(2026, 10, 12)))
        self.assertEqual(doc.charged, 2)
        with self.assertRaises(ValidationError):
            self.leave(start_date=date(2026, 10, 17), end_date=date(2026, 10, 18))

    def test_time_overlap_blocks_but_touching_endpoints_allowed(self):
        fields = dict(leave_type='outing', start_date=date(2026, 10, 12), end_date=date(2026, 10, 12))
        first = self.leave(**fields, start_time=time(9), end_time=time(12, 30))
        with self.assertRaises(ValidationError):
            self.leave(**fields, start_time=time(12), end_time=time(13))
        second = self.leave(**fields, start_time=time(12, 30), end_time=time(14))
        self.assertNotEqual(first.pk, second.pk)
        with self.assertRaises(ValidationError):
            self.leave()

    def test_draft_cancelled_rejected_and_deleted_do_not_block_overlap(self):
        self.leave(submit=False)
        for action in ['reject', 'delete', 'cancel_review']:
            doc = self.leave()
            if action == 'cancel_review':
                act_on_document(doc.pk, self.reviewer, 'review')
            act_on_document(doc.pk, self.owner if action == 'delete' else self.reviewer, action, '일정 변경')
        self.leave()

    def test_self_review_and_approval_still_require_two_actions(self):
        doc = self.leave(reviewer=self.owner, approver=self.owner)
        with self.assertRaises(PermissionDenied):
            act_on_document(doc.pk, self.owner, 'approve')
        act_on_document(doc.pk, self.owner, 'review')
        self.assertEqual(self.used(), 3)
        act_on_document(doc.pk, self.owner, 'approve')
        self.assertEqual(self.used(), 6)

    def test_submitted_document_is_immutable_and_delete_closes_after_review(self):
        doc = self.leave()
        doc.reason = '수정 요청'
        with self.assertRaises(PermissionDenied):
            save_document(doc, self.owner)
        doc.refresh_from_db()
        self.assertEqual(doc.reason, '개인 일정')
        act_on_document(doc.pk, self.reviewer, 'review')
        with self.assertRaises(PermissionDenied):
            act_on_document(doc.pk, self.owner, 'delete')
        self.client.force_login(self.owner)
        self.assertEqual(self.client.get(reverse('edit', args=[doc.pk])).status_code, 404)

    def test_actor_stage_permissions_and_required_reasons(self):
        doc = self.leave()
        for person, action in [(self.stranger, 'review'), (self.admin, 'review'), (self.approver, 'reject')]:
            with self.assertRaises(PermissionDenied):
                act_on_document(doc.pk, person, action, '사유')
        with self.assertRaises(ValidationError):
            act_on_document(doc.pk, self.reviewer, 'reject')
        act_on_document(doc.pk, self.reviewer, 'review')
        act_on_document(doc.pk, self.reviewer, 'cancel_review', '계획 변경')
        doc.refresh_from_db()
        self.assertEqual(doc.status, 'cancelled')
        self.assertEqual(self.used(), 3)

    def test_reassign_pending_stage_preserves_review_history(self):
        doc = self.leave()
        act_on_document(doc.pk, self.reviewer, 'review')
        doc = act_on_document(doc.pk, self.admin, 'reassign', '담당자 부재', self.stranger)
        self.assertEqual(doc.reviewer, self.reviewer)
        self.assertEqual(doc.approver, self.stranger)
        self.assertTrue(doc.reviewed_at)
        with self.assertRaises(PermissionDenied):
            act_on_document(doc.pk, self.approver, 'approve')
        act_on_document(doc.pk, self.stranger, 'approve')

    def test_accounting_visibility_excludes_cancelled_even_via_notice(self):
        doc = self.leave()
        self.assertFalse(visible_documents(self.accounting).filter(pk=doc.pk).exists())
        self.finish(doc)
        self.assertTrue(visible_documents(self.accounting).filter(pk=doc.pk).exists())
        act_on_document(doc.pk, self.approver, 'cancel_approval', '취소')
        self.assertFalse(visible_documents(self.accounting).filter(pk=doc.pk).exists())
        self.client.force_login(self.accounting)
        self.assertEqual(self.client.get(reverse('detail', args=[doc.pk])).status_code, 404)
        notice = Notice.objects.filter(user=self.accounting, document=doc).first()
        response = self.client.post(reverse('open_notice', args=[notice.pk]))
        self.assertRedirects(response, reverse('notices'))
        self.assertTrue(visible_documents(self.admin).filter(pk=doc.pk).exists())

    def test_purchase_total_shipping_recipient_and_cancel_freeze(self):
        doc = self.purchase()
        self.assertEqual(doc.total_price, Decimal('26001'))
        with self.assertRaises(PermissionDenied):
            act_on_document(doc.pk, self.accounting, 'place')
        self.finish(doc)
        with self.assertRaises(PermissionDenied):
            act_on_document(doc.pk, self.owner, 'place')
        act_on_document(doc.pk, self.accounting, 'place')
        with self.assertRaises(PermissionDenied):
            act_on_document(doc.pk, self.owner, 'receive')
        doc = act_on_document(doc.pk, self.recipient, 'receive')
        self.assertEqual(doc.shipment, 'received')
        second = self.finish(self.purchase(kind='stock', url=''))
        act_on_document(second.pk, self.accounting, 'place')
        act_on_document(second.pk, self.reviewer, 'cancel_review', '발주 취소')
        with self.assertRaises(PermissionDenied):
            act_on_document(second.pk, self.recipient, 'receive')

    def test_rejected_copy_is_new_editable_draft_with_fresh_number(self):
        doc = self.leave()
        act_on_document(doc.pk, self.reviewer, 'reject', '일정 확인')
        self.client.force_login(self.owner)
        response = self.client.post(reverse('copy', args=[doc.pk]))
        copy = Document.objects.filter(owner=self.owner, status='draft').get()
        self.assertRedirects(response, reverse('edit', args=[copy.pk]))
        self.assertIsNone(copy.number)
        save_document(copy, self.owner, submit=True)
        self.assertNotEqual(copy.number, doc.number)
        self.assertEqual(copy.charged, 0)

    def test_manual_remaining_edit_and_private_history(self):
        edit_balance(self.accounting, self.owner, 2026, 'remaining', Decimal('-2'), '긴급 사용분 반영')
        b = AnnualBalance.objects.get(user=self.owner, year=2026)
        self.assertEqual((b.total, b.used, b.remaining), (15, 17, -2))
        with self.assertRaises(PermissionDenied):
            edit_balance(self.admin, self.owner, 2026, 'total', Decimal(20), '수정')
        self.client.force_login(self.stranger)
        response = self.client.get(reverse('balances'))
        self.assertNotContains(response, '긴급 사용분 반영')
        self.client.force_login(self.owner)
        self.assertContains(self.client.get(reverse('balances')), '긴급 사용분 반영')

    def test_role_menus_and_server_authorization(self):
        self.client.force_login(self.owner)
        response = self.client.get(reverse('dashboard'))
        self.assertNotContains(response, '직원별 연차 관리')
        self.assertNotContains(response, '최종 승인 문서함')
        for url in [reverse('staff'), reverse('backups'), '/documents/?mode=all', '/documents/?mode=accounting']:
            self.assertEqual(self.client.get(url).status_code, 403)
        self.client.force_login(self.accounting)
        response = self.client.get(reverse('dashboard'))
        self.assertContains(response, '직원별 연차 관리')
        self.assertContains(response, '최종 승인 문서함')
        self.assertNotContains(response, '직원·권한 관리')

    def test_signup_approval_and_forced_password_change(self):
        response = self.client.post(reverse('signup'), {'first_name': '신규 직원', 'username': 'new-staff',
            'password1': 'staff-safe-password', 'password2': 'staff-safe-password', 'department': '생산팀', 'rank': '사원'})
        self.assertRedirects(response, reverse('login'))
        new = get_user_model().objects.get(username='new-staff')
        self.client.force_login(new)
        self.assertRedirects(self.client.get('/'), reverse('pending'))
        new.profile.approved = True
        new.profile.must_change_password = True
        new.profile.save()
        self.assertRedirects(self.client.get('/'), reverse('password'))
        response = self.client.post(reverse('password'), {'old_password': 'staff-safe-password',
            'new_password1': 'another-safe-password', 'new_password2': 'another-safe-password'})
        self.assertRedirects(response, '/')
        new.profile.refresh_from_db()
        self.assertFalse(new.profile.must_change_password)

    def test_signup_accepts_four_numeric_characters_and_rejects_short_password(self):
        data = {'first_name': '신규 직원', 'username': 'numeric-staff', 'password1': '123',
                'password2': '123', 'department': '생산팀', 'rank': '사원'}
        response = self.client.post(reverse('signup'), data)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(get_user_model().objects.filter(username='numeric-staff').exists())
        data['password1'] = data['password2'] = '1234'
        self.assertRedirects(self.client.post(reverse('signup'), data), reverse('login'))
        new = get_user_model().objects.get(username='numeric-staff')
        self.assertTrue(new.check_password('1234'))
        self.assertNotEqual(new.password, '1234')
        self.assertFalse(new.profile.department_confirmed)

    def test_employee_changes_only_own_password_and_stays_logged_in(self):
        self.client.force_login(self.owner)
        response = self.client.get(reverse('account'))
        self.assertContains(response, '내 계정')
        data = {'old_password': 'wrong', 'new_password1': '5678', 'new_password2': '5678',
                'user': self.stranger.pk}
        response = self.client.post(reverse('account'), data)
        self.assertContains(response, '현재 비밀번호')
        self.owner.refresh_from_db()
        self.assertTrue(self.owner.check_password('test-password-only'))
        data['old_password'] = 'test-password-only'
        data['new_password1'] = data['new_password2'] = '123'
        self.assertEqual(self.client.post(reverse('account'), data).status_code, 200)
        data['new_password1'], data['new_password2'] = '5678', '9999'
        self.assertEqual(self.client.post(reverse('account'), data).status_code, 200)
        data['new_password2'] = '5678'
        self.assertRedirects(self.client.post(reverse('account'), data), reverse('account'))
        self.owner.refresh_from_db()
        self.stranger.refresh_from_db()
        self.assertTrue(self.owner.check_password('5678'))
        self.assertTrue(self.stranger.check_password('test-password-only'))
        self.assertEqual(int(self.client.session['_auth_user_id']), self.owner.pk)
        self.assertTrue(Audit.objects.filter(actor=self.owner, event='본인 비밀번호 변경').exists())
        self.assertFalse(Audit.objects.filter(detail__contains='5678').exists())

    def test_forced_password_change_allows_four_numeric_characters(self):
        self.owner.profile.must_change_password = True
        self.owner.profile.save()
        self.client.force_login(self.owner)
        self.assertRedirects(self.client.get(reverse('account')), reverse('password'))
        self.assertRedirects(self.client.post(reverse('password'), {
            'old_password': 'test-password-only', 'new_password1': '4567', 'new_password2': '4567'}), '/')
        self.owner.refresh_from_db()
        self.assertTrue(self.owner.check_password('4567'))
        self.assertFalse(self.owner.profile.must_change_password)

    def test_staff_groups_new_signups_then_moves_them_to_assigned_department(self):
        self.client.post(reverse('signup'), {'first_name': '신규 생산직원', 'username': 'pending-staff',
            'password1': '1234', 'password2': '1234', 'department': '생산팀', 'rank': '사원'})
        new = get_user_model().objects.get(username='pending-staff')
        self.client.force_login(self.admin)
        response = self.client.get(reverse('staff'))
        self.assertEqual(response.context['selected_group'], 'unclassified')
        self.assertEqual([u.pk for u in response.context['staff']], [new.pk])
        response = self.client.get(reverse('staff'), {'department': '생산팀'})
        self.assertEqual(response.context['staff'], [])
        data = {'user': new.pk, 'department': '품질팀', 'rank': '주임', 'approved': 'on', 'active': 'on'}
        self.assertRedirects(self.client.post(reverse('staff') + '?department=unclassified', data),
            reverse('staff') + '?' + urlencode({'department': '품질팀'}))
        new.profile.refresh_from_db()
        self.assertTrue(new.profile.department_confirmed)
        response = self.client.get(reverse('staff'), {'department': '품질팀'})
        self.assertEqual([u.pk for u in response.context['staff']], [new.pk])
        self.assertEqual(response.context['staff_total'], 8)
        data.pop('approved')
        self.client.post(reverse('staff'), data)
        self.assertEqual([u.pk for u in self.client.get(reverse('staff'), {'department': '품질팀'}).context['staff']], [new.pk])
        self.assertEqual(self.client.get(reverse('staff'), {'department': 'unclassified'}).context['staff'], [])

    def test_staff_department_tabs_limit_results_and_permission(self):
        self.owner.profile.department = '생산팀'
        self.owner.profile.save()
        self.client.force_login(self.admin)
        response = self.client.get(reverse('staff'), {'department': '생산팀'})
        self.assertEqual([u.pk for u in response.context['staff']], [self.owner.pk])
        response = self.client.get(reverse('staff'), {'department': '영업팀'})
        self.assertContains(response, '이 부서에 등록된 직원이 없습니다.')
        self.assertEqual(self.client.get(reverse('staff'), {'department': 'invalid'}).status_code, 200)
        self.client.force_login(self.owner)
        self.assertEqual(self.client.get(reverse('staff'), {'department': '생산팀'}).status_code, 403)

    def test_approval_candidates_exclude_admin_and_juniors_but_allow_senior_self(self):
        self.recipient.profile.rank = '사원'
        self.recipient.profile.save()
        form = DocumentForm()
        for field in ['reviewer', 'approver']:
            self.assertFalse(form.fields[field].queryset.filter(pk=self.admin.pk).exists())
            self.assertFalse(form.fields[field].queryset.filter(pk=self.recipient.pk).exists())
            self.assertTrue(form.fields[field].queryset.filter(pk=self.owner.pk).exists())
        self.assertTrue(form.fields['recipient'].queryset.filter(pk=self.recipient.pk).exists())
        self.assertTrue(form.fields['recipient'].queryset.filter(pk=self.admin.pk).exists())
        self.stranger.profile.rank = '주임'
        self.stranger.profile.save()
        self.assertTrue(DocumentForm().fields['reviewer'].queryset.filter(pk=self.stranger.pk).exists())
        User = get_user_model()
        named_admin = User.objects.create_user('ADMIN', first_name='관리용 계정')
        Profile.objects.create(user=named_admin, department='인사팀', rank='이사', approved=True)
        self.assertFalse(DocumentForm().fields['approver'].queryset.filter(pk=named_admin.pk).exists())

    def test_direct_submission_and_reassignment_cannot_bypass_candidate_rules(self):
        self.stranger.profile.rank = '사원'
        self.stranger.profile.save()
        for candidate in [self.admin, self.stranger]:
            for field in ['reviewer', 'approver']:
                with self.subTest(candidate=candidate.username, field=field):
                    with self.assertRaises(ValidationError):
                        self.leave(**{field: candidate})
        doc = self.leave()
        for candidate in [self.admin, self.stranger]:
            with self.assertRaises(ValidationError):
                act_on_document(doc.pk, self.admin, 'reassign', '담당자 변경', candidate)
        doc.refresh_from_db()
        self.assertEqual(doc.reviewer, self.reviewer)
        self.client.force_login(self.owner)
        response = self.client.post(reverse('compose'), {'kind': 'stock', 'intent': 'submit',
            'product': '매뉴얼', 'quantity': 10, 'reviewer': self.admin.pk})
        self.assertEqual(response.status_code, 200)
        self.assertIn('reviewer', response.context['form'].errors)

    def test_department_picker_preserves_draft_and_rejects_mismatched_staff(self):
        self.approver.profile.department = '인사팀'
        self.approver.profile.save()
        draft = self.leave(submit=False)
        form = DocumentForm(instance=draft)
        self.assertEqual(form['reviewer_department'].value(), '개발팀')
        self.assertEqual(form['approver_department'].value(), '인사팀')
        self.assertIn('data-department="인사팀"', str(form['approver']))
        self.client.force_login(self.owner)
        data = {'kind': 'stock', 'intent': 'submit', 'product': '매뉴얼', 'quantity': 10,
                'reviewer': self.reviewer.pk, 'reviewer_department': '인사팀'}
        response = self.client.post(reverse('compose'), data)
        self.assertIn('reviewer', response.context['form'].errors)
        data['reviewer_department'] = '개발팀'
        response = self.client.post(reverse('compose'), data)
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Document.objects.filter(kind='stock', status='review').count(), 1)

    def test_copy_old_rejected_document_clears_now_ineligible_approval_people(self):
        old = self.leave()
        act_on_document(old.pk, self.reviewer, 'reject', '재검토')
        self.reviewer.profile.rank = '사원'
        self.reviewer.profile.save()
        self.client.force_login(self.owner)
        response = self.client.post(reverse('copy', args=[old.pk]))
        self.assertEqual(response.status_code, 302)
        copy = Document.objects.filter(status='draft').get()
        self.assertIsNone(copy.reviewer)
        old.refresh_from_db()
        self.assertEqual(old.reviewer, self.reviewer)

    def test_accounting_tabs_separate_approved_kinds_and_keep_filters(self):
        leave = self.finish(self.leave())
        office = self.finish(self.purchase())
        stock = self.finish(self.purchase(kind='stock', product='제품 매뉴얼'))
        self.purchase(product='검토 전 구매')
        self.client.force_login(self.accounting)
        response = self.client.get(reverse('documents'), {'mode': 'accounting'})
        self.assertEqual(response.context['selected_kind'], 'all')
        self.assertEqual({d.pk for d in response.context['documents']}, {leave.pk, office.pk, stock.pk})
        self.assertNotContains(response, '최종승인대기</a>')
        self.assertEqual({c['key']: c['count'] for c in response.context['accounting_categories']},
                         {'all': 3, 'leave': 1, 'office': 1, 'stock': 1})
        for kind, expected in [('office', office), ('stock', stock)]:
            response = self.client.get(reverse('documents'), {'mode': 'accounting', 'kind': kind})
            self.assertEqual([d.pk for d in response.context['documents']], [expected.pk])
        response = self.client.get(reverse('documents'), {'mode': 'accounting', 'kind': 'office', 'q': 'A4'})
        self.assertContains(response, 'name="kind" value="office"')
        self.assertTrue(all('q=A4' in c['url'] for c in response.context['accounting_categories']))
        act_on_document(office.pk, self.approver, 'cancel_approval', '취소')
        response = self.client.get(reverse('documents'), {'mode': 'accounting', 'kind': 'office'})
        self.assertEqual(list(response.context['documents']), [])
        self.client.force_login(self.owner)
        self.assertEqual(self.client.get(reverse('documents'), {'mode': 'accounting', 'kind': 'stock'}).status_code, 403)

    def test_accounting_unread_tabs_acknowledge_each_kind_and_all_per_user(self):
        self.finish(self.leave())
        self.finish(self.purchase())
        self.finish(self.purchase(kind='stock'))
        self.client.force_login(self.accounting)
        def states():
            return {tab['key']: tab['unread'] for tab in self.client.get(reverse('accounting_tab_status')).json()['tabs']}
        response = self.client.get(reverse('documents'), {'mode': 'accounting'})
        self.assertTrue(all(tab['unread'] for tab in response.context['accounting_categories']))
        self.assertFalse(AccountingReadState.objects.exists())
        self.assertTrue(all(states().values()))
        self.assertFalse(AccountingReadState.objects.exists())  # Polling never clears indicators.
        self.client.get(reverse('documents'), {'mode': 'accounting', 'kind': 'office'})
        self.assertEqual(states(), {'all': True, 'leave': True, 'office': False, 'stock': True})
        self.client.logout()
        self.client.force_login(self.accounting)
        self.assertFalse(states()['office'])  # Retained across sessions.
        self.owner.profile.view_accounting = True
        self.owner.profile.save()
        self.client.force_login(self.owner)
        self.assertTrue(states()['office'])  # Independent for each accountant.
        self.client.force_login(self.accounting)
        self.client.get(reverse('documents'), {'mode': 'accounting', 'kind': 'all'})
        self.assertFalse(any(states().values()))
        self.client.force_login(self.stranger)
        self.assertEqual(self.client.get(reverse('accounting_tab_status')).status_code, 403)

    def test_accounting_new_final_approvals_detected_and_cancelled_not_counted(self):
        pending_office = self.purchase()
        pending_stock = self.purchase(kind='stock')
        self.client.force_login(self.accounting)
        self.client.get(reverse('documents'), {'mode': 'accounting', 'kind': 'all'})
        self.finish(pending_office)
        self.finish(pending_stock)
        tabs = {tab['key']: tab for tab in self.client.get(reverse('accounting_tab_status')).json()['tabs']}
        self.assertTrue(tabs['all']['unread'])
        self.assertTrue(tabs['office']['unread'])
        self.assertTrue(tabs['stock']['unread'])
        self.assertFalse(tabs['leave']['unread'])
        self.assertEqual(tabs['all']['count'], 2)
        act_on_document(pending_office.pk, self.approver, 'cancel_approval', '취소')
        tabs = {tab['key']: tab for tab in self.client.get(reverse('accounting_tab_status')).json()['tabs']}
        self.assertEqual(tabs['office']['count'], 0)
        self.assertFalse(tabs['office']['unread'])
        self.assertTrue(tabs['all']['unread'])
        response = self.client.get(reverse('documents'), {'mode': 'accounting', 'kind': 'stock'})
        self.assertFalse(any(tab['unread'] for tab in response.context['accounting_categories']))

    def test_document_department_dropdown_removed_without_losing_leave_filter(self):
        doc = self.finish(self.purchase())
        self.client.force_login(self.accounting)
        response = self.client.get(reverse('documents'), {'mode': 'accounting', 'department': '품질팀'})
        self.assertNotContains(response, 'name="department"')
        self.assertEqual([d.pk for d in response.context['documents']], [doc.pk])
        self.assertTrue(all('department=' not in tab['url'] for tab in response.context['accounting_categories']))
        self.assertContains(self.client.get(reverse('balances')), 'name="department"')

    def test_csrf_is_enforced_for_mutation(self):
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.owner)
        doc = self.leave()
        self.assertEqual(client.post(reverse('action', args=[doc.pk]), {'action': 'delete'}).status_code, 403)

    def test_all_main_templates_render_for_each_role(self):
        doc = self.leave()
        for user in [self.owner, self.accounting, self.admin]:
            self.client.force_login(user)
            for url in ['/', '/compose/', '/compose/?kind=office', '/compose/?kind=stock', '/documents/',
                        '/documents/?mode=pending', '/documents/?mode=archive', '/documents/?mode=stock', '/balances/', '/policy/', '/notices/']:
                with self.subTest(user=user.username, url=url):
                    self.assertEqual(self.client.get(url).status_code, 200)
        self.assertEqual(self.client.get(reverse('detail', args=[doc.pk])).status_code, 200)
        self.assertEqual(self.client.get(reverse('staff')).status_code, 200)
        self.assertEqual(self.client.get(reverse('backups')).status_code, 200)

    def test_stock_requires_only_reviewer_and_has_no_money_or_final_approval(self):
        doc = self.purchase(kind='stock', approver=None, unit_price=0, url='', quantity=50)
        self.assertIsNone(doc.approver)
        self.assertEqual(doc.unit_price, 0)
        self.assertEqual(doc.status, 'review')
        with self.assertRaises(PermissionDenied):
            act_on_document(doc.pk, self.approver, 'approve')
        doc = act_on_document(doc.pk, self.reviewer, 'review')
        self.assertEqual(doc.status_label, '발주대기')
        self.assertEqual(self.used(), 3)
        self.assertTrue(Notice.objects.filter(user=self.accounting, document=doc, text__contains='발주 담당자').exists())
        with self.assertRaises(PermissionDenied):
            act_on_document(doc.pk, self.reviewer, 'review')
        self.client.force_login(self.owner)
        response = self.client.get(reverse('detail', args=[doc.pk]))
        self.assertContains(response, '요청 수량')
        self.assertNotContains(response, '수량 / 단가')
        self.assertNotContains(response, '총가격')
        self.assertNotContains(response, '03 최종 승인')

    def test_procurement_only_role_receives_stock_requests_without_accounting_access(self):
        self.stranger.profile.procure = True
        self.stranger.profile.save()
        stock = self.purchase(kind='stock', approver=None, unit_price=0, url='')
        office = self.finish(self.purchase())
        leave = self.finish(self.leave())
        self.assertFalse(visible_documents(self.stranger).filter(pk=stock.pk).exists())
        act_on_document(stock.pk, self.reviewer, 'review')
        self.assertTrue(visible_documents(self.stranger).filter(pk=stock.pk).exists())
        self.assertFalse(visible_documents(self.stranger).filter(pk__in=[office.pk, leave.pk]).exists())
        self.client.force_login(self.stranger)
        response = self.client.get('/documents/?mode=stock&stock_stage=ready')
        self.assertContains(response, stock.number)
        self.assertNotContains(response, office.number)
        self.assertNotContains(response, '최종승인대기')
        self.assertEqual(self.client.get('/documents/?mode=accounting').status_code, 403)
        response = self.client.post(reverse('action', args=[stock.pk]), {'action': 'place'})
        self.assertRedirects(response, reverse('detail', args=[stock.pk]))
        stock.refresh_from_db()
        self.assertEqual(stock.status_label, '입고대기')
        self.assertNotContains(self.client.get('/documents/?mode=stock&stock_stage=ready'), stock.number)
        self.assertContains(self.client.get('/documents/?mode=stock&stock_stage=ordered'), stock.number)

    def test_stock_review_cancellation_blocks_fulfillment_and_requires_own_reviewer(self):
        doc = self.finish(self.purchase(kind='stock', approver=None, unit_price=0, url=''))
        act_on_document(doc.pk, self.accounting, 'place')
        for user in [self.owner, self.accounting, self.approver]:
            with self.assertRaises(PermissionDenied):
                act_on_document(doc.pk, user, 'cancel_review', '취소 사유')
        act_on_document(doc.pk, self.reviewer, 'cancel_review', '생산 계획 변경')
        self.assertFalse(visible_documents(self.accounting).filter(pk=doc.pk).exists())
        with self.assertRaises(PermissionDenied):
            act_on_document(doc.pk, self.recipient, 'receive')
        with self.assertRaises(PermissionDenied):
            act_on_document(doc.pk, self.accounting, 'place')
        self.assertEqual(self.used(), 3)

    def test_stock_form_submission_ignores_price_and_approver_and_can_self_review(self):
        self.client.force_login(self.owner)
        response = self.client.post(reverse('compose'), {'intent': 'submit', 'kind': 'stock',
            'product': '제품 보관함', 'quantity': 20, 'reviewer': self.owner.pk,
            'unit_price': '12345', 'approver': self.approver.pk, 'url': 'https://example.com/unused'})
        doc = Document.objects.get(kind='stock')
        self.assertRedirects(response, reverse('detail', args=[doc.pk]))
        self.assertIsNone(doc.approver)
        self.assertEqual(doc.unit_price, 0)
        self.assertEqual(doc.url, '')
        reviewed = act_on_document(doc.pk, self.owner, 'review')
        self.assertEqual(reviewed.status_label, '발주대기')
        self.assertContains(self.client.get('/documents/?mode=history'), doc.number)

    def test_stock_reject_copy_retains_single_review_workflow(self):
        doc = self.purchase(kind='stock', approver=None, unit_price=0, url='')
        act_on_document(doc.pk, self.reviewer, 'reject', '수량 확인')
        self.client.force_login(self.owner)
        self.client.post(reverse('copy', args=[doc.pk]))
        copy = Document.objects.get(status='draft', kind='stock')
        self.assertIsNone(copy.approver)
        self.assertEqual(copy.unit_price, 0)
        save_document(copy, self.owner, submit=True)
        self.assertNotEqual(copy.number, doc.number)
        self.assertEqual(act_on_document(copy.pk, self.reviewer, 'review').status, 'approved')

    def test_legacy_stock_conversion_preserves_existing_reviews_and_fulfillment(self):
        import importlib
        from types import SimpleNamespace
        from django.apps import apps
        from django.db import connection
        from django.utils import timezone
        pending = Document.objects.create(owner=self.owner, reviewer=self.reviewer, approver=self.approver,
            kind='stock', status='approve', number='LEGACY-1', title='구매요청 / owner / 보관함',
            product='보관함', quantity=20, unit_price=15000, reviewed_at=timezone.now())
        old_log = Audit.objects.create(document=pending, actor=self.reviewer, event='검토 승인')
        ordered = Document.objects.create(owner=self.owner, reviewer=self.reviewer, approver=self.approver,
            kind='stock', status='approved', shipment='ordered', number='LEGACY-2', product='매뉴얼', quantity=100)
        unreviewed = Document.objects.create(owner=self.owner, reviewer=self.reviewer, kind='stock', status='approve')
        migration = importlib.import_module('workflow.migrations.0003_stock_review_workflow')
        migration.convert_stock_requests(apps, SimpleNamespace(connection=connection))
        pending.refresh_from_db(); ordered.refresh_from_db(); unreviewed.refresh_from_db()
        self.assertEqual(pending.status_label, '발주대기')
        self.assertEqual(pending.approved_at, pending.reviewed_at)
        self.assertEqual((pending.number, pending.quantity, pending.unit_price), ('LEGACY-1', 20, 15000))
        self.assertTrue(Audit.objects.filter(pk=old_log.pk).exists())
        self.assertTrue(Audit.objects.filter(document=pending, actor__isnull=True).exists())
        self.assertEqual(ordered.shipment, 'ordered')
        self.assertEqual(unreviewed.status, 'review')
        self.assertTrue(Notice.objects.filter(document=pending, user=self.accounting).exists())
